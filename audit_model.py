"""
Independent audit of Joby_Financial_Model.xlsx.

For each test case the script copies the workbook, edits inputs with openpyxl,
recalculates the copy in headless LibreOffice (NOT Microsoft Excel), and then:

  * recomputes every monthly and annual line in plain Python from the raw
    inputs and compares it cell by cell with the workbook's formula results
    (Financial Model tab, the Scenario Comparison engine for all three cases,
    and the Scenario Comparison summary table);
  * tests the accounting identities directly on the workbook's own values:
        revenue            = passengers flown x average fare
        passengers flown   = completed flights x seats per flight x load factor
        contribution       = revenue - direct flight costs - passenger-variable costs
        route op. profit   = contribution - fixed route costs
        annual totals      = sum of months; annual rates = ratio of annual totals
        breakeven LF       -> profit (contribution or route) is exactly zero at that LF
  * confirms ratios with a zero denominator show "n.m." instead of a number,
    and that the integrity checks give the expected result.

Usage:
    python audit_model.py [Joby_Financial_Model.xlsx] [--report audit_report.md]
"""

import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

from validate_model import LibreOffice, scan_errors

REL = 1e-9
ABS = 1e-6
NM = "n.m."
MONTH_COLS = list(range(3, 15))  # C..N
TOTAL = 15  # O
SCEN_COL = {"Base": 3, "Upside": 4, "Downside": 5}

ASSUMPTION_LABELS = {
    "capacity": "Aircraft passenger seats per flight",
    "fpd": "Scheduled flights per operating day",
    "days": "Operating days per month",
    "completion": "Flight completion rate",
    "fare": "Average fare per passenger (Month 1)",
    "lf": "Passenger load factor (Month 1)",
    "vcf": "Variable operating cost per completed flight (Month 1)",
    "vcp": "Variable cost per passenger (Month 1)",
    "fixed": "Monthly route-specific fixed costs (Month 1)",
    "fg": "Monthly fare growth rate",
    "dg": "Monthly demand growth rate",
    "cg": "Monthly operating cost growth rate",
}
GENERAL = {"capacity", "fpd", "days"}

# Financial Model rows (exact labels) -> reference key
FM_LINES = {
    "Operating days": "days",
    "Scheduled flights per operating day": "fpd",
    "Flight completion rate": "comp",
    "Passenger load factor": "lf",
    "Variable operating cost per completed flight": "vcf",
    "Variable cost per passenger": "vcp",
    "Scheduled flights": "sched",
    "Completed flights": "completed",
    "Available passenger seats": "seats",
    "Passengers flown": "pax",
    "Average fare": "fare",
    "Passenger revenue": "revenue",
    "Direct flight operating costs": "flight_cost",
    "Passenger-variable costs": "pax_cost",
    "Total variable costs": "tvc",
    "Contribution profit (flight-level)": "contrib",
    "Contribution margin": "cm_pct",
    "Contribution profit per completed flight": "contrib_flt",
    "Fixed route costs": "fixed",
    "Route operating profit / (loss)": "rop",
    "Route operating margin": "rom",
    "Cumulative route operating profit / (loss)": "cum_rop",
    "Net margin per passenger (fare - variable cost per passenger)": "net_pax_margin",
    "Passenger margin if every seat were sold (100% load factor)": "max_margin",
    "Breakeven load factor - flight level (contribution = 0)": "be_flight",
    "Breakeven load factor - route (after fixed route costs)": "be_route",
    "Load factor headroom vs. route breakeven": "lf_headroom",
}
ENGINE_LINES = {
    "Average fare": "fare",
    "Passenger load factor": "lf",
    "Scheduled flights": "sched",
    "Completed flights": "completed",
    "Available passenger seats": "seats",
    "Passengers flown": "pax",
    "Passenger revenue": "revenue",
    "Direct flight operating costs": "flight_cost",
    "Passenger-variable costs": "pax_cost",
    "Total variable costs": "tvc",
    "Contribution profit": "contrib",
    "Fixed route costs": "fixed",
    "Route operating profit / (loss)": "rop",
    "Passenger margin at 100% load factor": "max_margin",
}
SUMMARY_LINES = {
    "Scheduled flights": "sched",
    "Completed flights": "completed",
    "Passengers flown": "pax",
    "Average load factor (weighted)": "lf",
    "Average fare (weighted)": "fare",
    "Annual passenger revenue": "revenue",
    "Total variable costs": "tvc",
    "Contribution profit": "contrib",
    "Contribution margin": "cm_pct",
    "Contribution profit per completed flight": "contrib_flt",
    "Fixed route costs": "fixed",
    "Route operating profit / (loss)": "rop",
    "Route operating margin": "rom",
    "Breakeven load factor - flight level": "be_flight",
    "Breakeven load factor - route": "be_route",
}


# --------------------------------------------------------------------------
# Independent reference model
# --------------------------------------------------------------------------
def div(a, b):
    return a / b if b else NM


def reference(p):
    """Recompute every line from raw inputs. Returns (monthly dict, annual dict)."""
    M = {k: [] for k in FM_LINES.values()}
    cum = 0.0
    for m in range(12):
        fare = p["fare"] * (1 + p["fg"]) ** m
        lf = min(1.0, max(0.0, p["lf"] * (1 + p["dg"]) ** m))
        ci = (1 + p["cg"]) ** m
        vcf, vcp = p["vcf"] * ci, p["vcp"] * ci
        sched = p["days"] * p["fpd"]
        completed = sched * p["completion"]
        seats = completed * p["capacity"]
        pax = completed * p["capacity"] * lf
        revenue = pax * fare
        flight_cost = completed * vcf
        pax_cost = pax * vcp
        tvc = flight_cost + pax_cost
        contrib = revenue - flight_cost - pax_cost
        fixed = p["fixed"] * ci
        rop = contrib - fixed
        cum += rop
        max_margin = seats * (fare - vcp)
        be_flight = flight_cost / max_margin if max_margin > 0 else NM
        be_route = (flight_cost + fixed) / max_margin if max_margin > 0 else NM
        vals = dict(days=p["days"], fpd=p["fpd"], comp=p["completion"], lf=lf, vcf=vcf, vcp=vcp,
                    sched=sched, completed=completed, seats=seats, pax=pax, fare=fare,
                    revenue=revenue, flight_cost=flight_cost, pax_cost=pax_cost, tvc=tvc,
                    contrib=contrib, cm_pct=div(contrib, revenue),
                    contrib_flt=div(contrib, completed), fixed=fixed, rop=rop,
                    rom=div(rop, revenue), cum_rop=cum, net_pax_margin=fare - vcp,
                    max_margin=max_margin, be_flight=be_flight, be_route=be_route,
                    lf_headroom=lf - be_route if be_route != NM else NM)
        for k, v in vals.items():
            M[k].append(v)

    S = {k: sum(M[k]) for k in ("days", "sched", "completed", "seats", "pax", "revenue",
                                "flight_cost", "pax_cost", "tvc", "contrib", "fixed", "rop",
                                "max_margin")}
    Y = dict(S)
    Y.update(
        fpd=div(S["sched"], S["days"]),
        comp=div(S["completed"], S["sched"]),
        lf=div(S["pax"], S["seats"]),
        vcf=div(S["flight_cost"], S["completed"]),
        vcp=div(S["pax_cost"], S["pax"]),
        fare=div(S["revenue"], S["pax"]),
        cm_pct=div(S["contrib"], S["revenue"]),
        contrib_flt=div(S["contrib"], S["completed"]),
        rom=div(S["rop"], S["revenue"]),
        cum_rop=M["cum_rop"][-1],
        net_pax_margin=div(S["revenue"] - S["pax_cost"], S["pax"]),
        be_flight=S["flight_cost"] / S["max_margin"] if S["max_margin"] > 0 else NM,
        be_route=(S["flight_cost"] + S["fixed"]) / S["max_margin"] if S["max_margin"] > 0 else NM,
    )
    Y["lf_headroom"] = (Y["lf"] - Y["be_route"]
                        if NM not in (Y["lf"], Y["be_route"]) else NM)
    return M, Y


# --------------------------------------------------------------------------
# Workbook helpers
# --------------------------------------------------------------------------
def rows_by_label(ws, start=1, stop=None):
    out = {}
    for r in range(start, (stop or ws.max_row) + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and v not in out:
            out[v.strip()] = r
    return out


def same(wb_val, ref):
    if ref == NM:
        return wb_val == NM
    if not isinstance(wb_val, (int, float)) or isinstance(wb_val, bool):
        return False
    return abs(wb_val - ref) <= max(ABS, abs(ref) * REL)


def read_inputs(wb):
    ws = wb["Assumptions"]
    rows = rows_by_label(ws)
    out = {}
    for s, col in SCEN_COL.items():
        out[s] = {k: ws.cell(row=rows[lbl], column=3 if k in GENERAL else col).value
                  for k, lbl in ASSUMPTION_LABELS.items()}
    return out, rows


class Audit:
    def __init__(self):
        self.comparisons = 0
        self.identities = 0
        self.failures = []

    def cmp(self, where, wb_val, ref):
        self.comparisons += 1
        if not same(wb_val, ref):
            self.failures.append(f"{where}: workbook={wb_val!r} reference={ref!r}")

    def ident(self, where, ok, detail=""):
        self.identities += 1
        if not ok:
            self.failures.append(f"IDENTITY {where} {detail}")


def close(a, b):
    return abs(a - b) <= max(ABS, abs(b) * REL)


def check_identities(au, tag, V, Y, capacity, fixed_in_rows=True):
    """V: dict key -> list of 12 workbook values; Y: dict key -> annual workbook value."""
    for i in range(12):
        m = f"{tag} M{i + 1}"
        g = lambda k: V[k][i]  # noqa: E731
        au.ident(f"{m} revenue = pax x fare", close(g("revenue"), g("pax") * g("fare")))
        au.ident(f"{m} pax = completed x seats/flight x LF",
                 close(g("pax"), g("completed") * capacity * g("lf")))
        au.ident(f"{m} contribution = revenue - flight costs - pax costs",
                 close(g("contrib"), g("revenue") - g("flight_cost") - g("pax_cost")))
        au.ident(f"{m} route op. profit = contribution - fixed",
                 close(g("rop"), g("contrib") - g("fixed")))
        # Breakeven: profit at the breakeven LF must be exactly zero, with
        # passenger-variable costs scaling with passengers at that LF.
        if "be_route" in V and isinstance(g("be_route"), (int, float)):
            seats, fare = g("seats"), g("fare")
            vcp = g("vcp") if "vcp" in V else (g("pax_cost") / g("pax") if g("pax") else None)
            if vcp is not None:
                for key, fx in (("be_flight", 0.0), ("be_route", g("fixed"))):
                    lf = g(key)
                    profit = seats * lf * fare - g("flight_cost") - seats * lf * vcp - fx
                    au.ident(f"{m} profit at {key} = 0", abs(profit) < 1e-4, f"({profit:.6f})")
    for k in ("sched", "completed", "seats", "pax", "revenue", "flight_cost", "pax_cost", "tvc",
              "contrib", "fixed", "rop"):
        if k in V and k in Y:
            au.ident(f"{tag} annual {k} = sum of months", close(Y[k], sum(V[k])))
    au.ident(f"{tag} annual revenue = sum(pax_m x fare_m)",
             close(Y["revenue"], sum(p * f for p, f in zip(V["pax"], V["fare"]))))


# --------------------------------------------------------------------------
# Sensitivity Analysis audit
# --------------------------------------------------------------------------
SA_SETTINGS = {"fare_step": "Fare step between columns", "lf_start": "Load factor - first row",
               "lf_step": "Load factor step between rows", "band": "Breakeven band (+/-)",
               "cost_shock": "Illustrative direct flight cost change"}


def sensitivity_reference(p, st):
    cap, fare, lf, vcf, vcp, fixed = (p["capacity"], p["fare"], p["lf"], p["vcf"], p["vcp"],
                                      p["fixed"])
    completed = p["days"] * p["fpd"] * p["completion"]
    fares = [max(0.0, round(fare - 4 * st["fare_step"], 2))]
    for _ in range(8):
        fares.append(max(0.0, round(fares[-1] + st["fare_step"], 2)))
    lfs = [round(st["lf_start"], 4)]
    for _ in range(14):
        lfs.append(round(lfs[-1] + st["lf_step"], 4))
    grid = [[cap * l * (f - vcp) - vcf for f in fares] for l in lfs]
    fixed_pf = fixed / completed if completed > 0 else NM
    be_f = lambda f: vcf / (cap * (f - vcp)) if f > vcp else NM  # noqa: E731
    be_r = lambda f: ((vcf + fixed_pf) / (cap * (f - vcp))  # noqa: E731
                      if f > vcp and fixed_pf != NM else NM)
    bf, br = be_f(fare), be_r(fare)
    flat = [v for row in grid for v in row]
    band = st["band"]
    return dict(
        fares=fares, lfs=lfs, grid=grid, be_flight=bf, be_route=br, fixed_pf=fixed_pf,
        completed=completed,
        be_gap=br - bf if NM not in (bf, br) else NM,
        hr_flight=lf - bf if bf != NM else NM, hr_route=lf - br if br != NM else NM,
        be_flight_row=[be_f(f) for f in fares], be_route_row=[be_r(f) for f in fares],
        counts=[sum(v < -band for v in flat), sum(-band <= v <= band for v in flat),
                sum(v > band for v in flat)],
        lf_axis_ok=min(lfs) >= 0 and max(lfs) <= 1)


def audit_sensitivity(au, wb, p):
    ws = wb["Sensitivity Analysis"]
    rows = rows_by_label(ws)
    st = {k: ws.cell(row=rows[v], column=2).value for k, v in SA_SETTINGS.items()}
    ref = sensitivity_reference(p, st)
    lab = {
        "be_flight": "FLIGHT-LEVEL breakeven load factor (contribution per flight = $0)",
        "be_route": "ROUTE-LEVEL breakeven load factor (after fixed route costs)",
        "fixed_pf": "Fixed route costs per completed flight",
        "be_gap": "Gap: route-level minus flight-level breakeven",
        "hr_flight": "Base load factor headroom vs. flight-level breakeven",
        "hr_route": "Base load factor headroom vs. route-level breakeven",
        "completed": "Completed flights per month",
    }
    for k, l in lab.items():
        au.cmp(f"SA {k}", ws.cell(row=rows[l], column=2).value, ref[k])
    for k, l in (("capacity", "Aircraft passenger seats per flight"), ("fare", "Base average fare"),
                 ("lf", "Base passenger load factor"),
                 ("vcf", "Direct flight operating cost per completed flight"),
                 ("vcp", "Variable cost per passenger"), ("fixed", "Fixed route costs per month")):
        au.cmp(f"SA link {k}", ws.cell(row=rows[l], column=2).value, p[k])
    hdr = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=2).value == "LF \\ Fare")
    for j, f in enumerate(ref["fares"]):
        au.cmp(f"SA fare axis col{j + 1}", ws.cell(row=hdr, column=3 + j).value, f)
    for i, l in enumerate(ref["lfs"]):
        au.cmp(f"SA LF axis row{i + 1}", ws.cell(row=hdr + 1 + i, column=2).value, l)
        for j in range(9):
            v = ws.cell(row=hdr + 1 + i, column=3 + j).value
            au.cmp(f"SA grid LF={l} fare={ref['fares'][j]}", v, ref["grid"][i][j])
            # identity on the workbook's own axis values
            au.ident("SA grid identity", close(v, p["capacity"] * ws.cell(row=hdr + 1 + i, column=2).value
                                               * (ws.cell(row=hdr, column=3 + j).value - p["vcp"])
                                               - p["vcf"]))
    rf = rows["Flight-level breakeven LF at this fare"]
    rr = rows["Route-level breakeven LF at this fare (after fixed costs)"]
    for j in range(9):
        au.cmp(f"SA BE flight col{j + 1}", ws.cell(row=rf, column=3 + j).value, ref["be_flight_row"][j])
        au.cmp(f"SA BE route col{j + 1}", ws.cell(row=rr, column=3 + j).value, ref["be_route_row"][j])
        # contribution at the column's flight-level breakeven LF must be zero
        be = ref["be_flight_row"][j]
        if be != NM:
            au.ident("SA zero contribution at column BE",
                     abs(p["capacity"] * be * (ref["fares"][j] - p["vcp"]) - p["vcf"]) < 1e-6)
    for i, lbl in enumerate(("Loss-making (below -band)", "Breakeven (within +/- band)",
                             "Profitable (above +band)")):
        au.cmp(f"SA count {lbl}", ws.cell(row=rows[lbl], column=2).value, ref["counts"][i])
    # route-level breakeven really is route breakeven: monthly route profit = 0
    if ref["be_route"] != NM:
        prof = ref["completed"] * (p["capacity"] * ref["be_route"] * (p["fare"] - p["vcp"])
                                   - p["vcf"]) - p["fixed"]
        au.ident("SA route profit at route BE = 0", abs(prof) < 1e-6)
    for r in range(rows["5. Business Insights (illustrative; text updates with the inputs)"] + 1,
                   rows["6. Sensitivity Validation Checks"]):
        t = ws.cell(row=r, column=1).value
        if t:
            au.ident(f"SA insight row {r} is text without errors",
                     isinstance(t, str) and not any(e in t for e in ("#VALUE", "#DIV", "#NAME", "#REF")))
    first = rows["6. Sensitivity Validation Checks"] + 1
    statuses = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
                for r in range(first, rows["SENSITIVITY STATUS"])}
    return ref, statuses, ws.cell(row=rows["SENSITIVITY STATUS"], column=2).value


# --------------------------------------------------------------------------
# Executive Dashboard audit
# --------------------------------------------------------------------------
DB_KPIS = {"PASSENGER REVENUE": "revenue", "PASSENGERS FLOWN": "pax",
           "COMPLETED FLIGHTS": "completed", "LOAD FACTOR": "lf",
           "CONTRIBUTION / FLIGHT": "contrib_flt", "ROUTE OPERATING PROFIT": "rop"}
DB_TABLE = {"Passenger revenue": "revenue", "Passengers flown": "pax",
            "Load factor (weighted)": "lf", "Average fare ($ / passenger)": "fare",
            "Contribution per flight": "contrib_flt", "Contribution profit": "contrib",
            "Route operating profit": "rop", "Route operating margin": "rom",
            "Route breakeven load factor": "be_route", "Profitable months (of 12)": "months_pos"}


def audit_dashboard(au, wb, inputs, scenario):
    ws = wb["Executive Dashboard"]
    Ya = reference(inputs[scenario])[1]
    # KPI cards: label in row 5, value directly below
    found = 0
    for col in range(2, 14):
        label = ws.cell(row=5, column=col).value
        if label in DB_KPIS:
            found += 1
            au.cmp(f"DB KPI {label}", ws.cell(row=6, column=col).value, Ya[DB_KPIS[label]])
    au.ident("DB has 6 KPI cards", found == 6)
    # Scenario table vs reference for every scenario
    for r in range(11, 25):
        label = ws.cell(row=r, column=9).value
        if label not in DB_TABLE:
            continue
        for col, s in zip((11, 12, 13), ("Base", "Upside", "Downside")):
            M, Y = reference(inputs[s])
            ref = (sum(v > 0 for v in M["rop"]) if DB_TABLE[label] == "months_pos"
                   else Y[DB_TABLE[label]])
            au.cmp(f"DB table {label} {s}", ws.cell(row=r, column=col).value, ref)
    # Key decisions: three evidence cells, text, no errors
    for col in (2, 6, 10):
        t = ws.cell(row=28, column=col).value
        au.ident(f"DB decision text col{col}", isinstance(t, str) and len(t) > 40
                 and not any(e in t for e in ("#VALUE", "#DIV", "#NAME", "#REF", "#N/A")))
    # Subtitle reflects the scenario shown
    au.ident("DB subtitle names scenario", str(ws["B2"].value).endswith(scenario))
    status = None
    for r in range(30, ws.max_row + 1):
        if ws.cell(row=r, column=2).value == "DASHBOARD STATUS":
            status = ws.cell(row=r, column=10).value
    return status


def audit_dashboard_chart(path):
    """Read the chart XML from the delivered file and return its series references."""
    import zipfile
    import re
    with zipfile.ZipFile(path) as z:
        charts = [n for n in z.namelist() if n.startswith("xl/charts/chart")]
        import html
        refs = ([html.unescape(r) for r in re.findall(r"<c:f>([^<]+)</c:f>",
                                                      z.read(charts[0]).decode())]
                if charts else [])
    return charts, refs


# --------------------------------------------------------------------------
# Test cases: (name, active scenario, {input key: value applied to that scenario})
# --------------------------------------------------------------------------
CASES = [
    ("Base (as delivered)", "Base", {}),
    ("Upside (selector)", "Upside", {}),
    ("Downside (selector)", "Downside", {}),
    ("Zero flights (0 flights/day)", "Base", {"fpd": 0}),
    ("Zero completion rate", "Base", {"completion": 0}),
    ("Completion rate 50%", "Base", {"completion": 0.5}),
    ("Completion rate 100%", "Base", {"completion": 1.0}),
    ("Zero load factor", "Base", {"lf": 0}),
    ("Full load factor, no growth", "Base", {"lf": 1.0, "dg": 0}),
    ("Full load factor + demand growth (cap)", "Base", {"lf": 1.0}),
    ("Higher flight op. cost ($600/flight)", "Base", {"vcf": 600}),
    ("Fare below cost per passenger ($10)", "Base", {"fare": 10}),
    ("Sensitivity: $5 fare step, LF 50%-85% by 2.5pt", "Base",
     {"sa:fare_step": 5, "sa:lf_start": 0.5, "sa:lf_step": 0.025}),
    ("Sensitivity: wider $40 fare step, $0 band", "Base", {"sa:fare_step": 40, "sa:band": 0}),
    ("Sensitivity: LF axis past 100% (should FAIL)", "Base",
     {"sa:lf_start": 0.5, "sa:lf_step": 0.05}),
]
EXPECTED_FAILS = {
    "Sensitivity: LF axis past 100% (should FAIL)": {"Load factor axis within 0%-100%",
                                                    "Sensitivity Analysis checks pass"},
}


def apply_case(wb, scenario, overrides, rows):
    ws = wb["Assumptions"]
    ws["C5"] = scenario
    sa = wb["Sensitivity Analysis"]
    sarows = rows_by_label(sa)
    for k, v in overrides.items():
        if k.startswith("sa:"):
            sa.cell(row=sarows[SA_SETTINGS[k[3:]]], column=2, value=v)
            continue
        col = 3 if k in GENERAL else SCEN_COL[scenario]
        ws.cell(row=rows[ASSUMPTION_LABELS[k]], column=col, value=v)


def run_case(lo, src, work, idx, name, scenario, overrides):
    au = Audit()
    path = work / f"case{idx}.xlsx"
    shutil.copy(src, path)
    wb = load_workbook(path)
    _, arows = read_inputs(wb)
    apply_case(wb, scenario, overrides, arows)
    wb.save(path)
    lo.recalc(path)

    errors = scan_errors(path)
    for e in errors:
        au.failures.append(f"formula error {e}")
    wb = load_workbook(path, data_only=True)
    inputs, _ = read_inputs(wb)

    # Financial Model: every month and annual cell vs reference
    fm = wb["Financial Model"]
    frows = rows_by_label(fm)
    M, Y = reference(inputs[scenario])
    V, YW = {}, {}
    for label, key in FM_LINES.items():
        r = frows[label]
        V[key] = [fm.cell(row=r, column=c).value for c in MONTH_COLS]
        YW[key] = fm.cell(row=r, column=TOTAL).value
        for i, c in enumerate(MONTH_COLS):
            au.cmp(f"FM {label} M{i + 1}", V[key][i], M[key][i])
        if key in Y:
            au.cmp(f"FM {label} FY", YW[key], Y[key])
    check_identities(au, "FM", V, YW, inputs[scenario]["capacity"])

    # Scenario Comparison engine (all three cases) and summary
    sc = wb["Scenario Comparison"]
    srows = rows_by_label(sc, 1, 45)
    for s, col in SCEN_COL.items():
        Ms, Ys = reference(inputs[s])
        hdr = rows_by_label(sc, 46)[f"{s} case"]
        erows = rows_by_label(sc, hdr + 1, hdr + 17)
        EV, EY = {}, {}
        for label, key in ENGINE_LINES.items():
            r = erows[label]
            EV[key] = [sc.cell(row=r, column=c).value for c in MONTH_COLS]
            EY[key] = sc.cell(row=r, column=TOTAL).value
            for i, c in enumerate(MONTH_COLS):
                au.cmp(f"SC engine {s} {label} M{i + 1}", EV[key][i], Ms[key][i])
            if key in Ys and key not in ("lf", "fare"):
                au.cmp(f"SC engine {s} {label} FY", EY[key], Ys[key])
        check_identities(au, f"SC {s}", EV, EY, inputs[s]["capacity"])
        for label, key in SUMMARY_LINES.items():
            au.cmp(f"SC summary {s} {label}", sc.cell(row=srows[label], column=col).value, Ys[key])

    # Weighted-vs-simple average evidence (annual rates must be ratio of totals)
    notes = []
    if Y["fare"] != NM:
        simple = sum(M["fare"]) / 12
        notes.append(f"FY fare weighted ${Y['fare']:.2f} vs simple avg ${simple:.2f}")

    # Integrity checks
    first = frows["Check"] + 1
    last = frows["OVERALL MODEL STATUS"]
    statuses = {fm.cell(row=r, column=1).value: fm.cell(row=r, column=3).value
                for r in range(first, last)}
    info = {}
    sref, sstat, sa_status = audit_sensitivity(au, wb, inputs["Base"])
    expected = EXPECTED_FAILS.get(name, set())
    fails = {k for k, v in statuses.items() if v != "PASS"}
    sfails = {k for k, v in sstat.items() if v == "FAIL"}
    if (fails | sfails) != expected:
        au.failures.append(f"check results {sorted(fails | sfails)} != expected {sorted(expected)}")
    # The sheet's own checks must agree with the Python verdicts
    au.ident("SA LF-axis check agrees with Python",
             (sstat["Load factor axis within 0%-100%"] == "PASS") == sref["lf_axis_ok"])
    info["sensitivity"] = sa_status
    info["dashboard"] = audit_dashboard(au, wb, inputs, scenario)
    if not str(info["dashboard"]).startswith("ALL"):
        au.failures.append(f"dashboard status: {info['dashboard']}")
    master = fm.cell(row=last, column=3).value
    for lbl in ("Months where demand exceeds seat capacity (load factor capped at 100%)",):
        if lbl in frows:
            info["capped months"] = fm.cell(row=frows[lbl], column=3).value
    return au, Y, master, notes, info


def fmt_money(v):
    return "n.m." if v == NM else f"${v:,.0f}"


def fmt_pct(v):
    return "n.m." if v == NM else f"{v:.1%}"


def main():
    argv = sys.argv[1:]
    report = None
    if "--report" in argv:
        i = argv.index("--report")
        report = Path(argv[i + 1])
        argv = argv[:i] + argv[i + 2:]
    src = Path(argv[0] if argv else Path(__file__).with_name("Joby_Financial_Model.xlsx"))
    lo = LibreOffice()
    work = Path(lo.tmp.name)
    lines = []
    total_cmp = total_id = total_fail = 0
    print("Recalculation engine: LibreOffice headless (Microsoft Excel was not used)\n")
    hdr = (f"| Case | Revenue | Contribution | Route op. profit | Contrib./flight | "
           f"BE LF flight | BE LF route | Cells compared | Identities | Mismatches | Status | "
           f"Sensitivity status | Dashboard status |")
    lines += [hdr, "|" + "---|" * 13]
    detail = []
    for i, (name, scen, ov) in enumerate(CASES):
        au, Y, master, notes, info = run_case(lo, src, work, i, name, scen, ov)
        total_cmp += au.comparisons
        total_id += au.identities
        total_fail += len(au.failures)
        cpf = "n.m." if Y["contrib_flt"] == NM else f"${Y['contrib_flt']:,.2f}"
        status = master + (f"; {info['capped months']} capped month(s)" if info.get(
            "capped months") else "")
        lines.append(f"| {name} | {fmt_money(Y['revenue'])} | {fmt_money(Y['contrib'])} | "
                     f"{fmt_money(Y['rop'])} | {cpf} | {fmt_pct(Y['be_flight'])} | "
                     f"{fmt_pct(Y['be_route'])} | {au.comparisons} | {au.identities} | "
                     f"{len(au.failures)} | {status} | {info.get('sensitivity')} | {info.get('dashboard')} |")
        print(f"{name:<42} compared={au.comparisons:<5} identities={au.identities:<5} "
              f"mismatches={len(au.failures):<4} {master}  {'; '.join(notes)}")
        for f in au.failures[:12]:
            print("    -", f)
            detail.append(f"- {name}: {f}")
        if len(au.failures) > 12:
            print(f"    ... {len(au.failures) - 12} more")
            detail.append(f"- {name}: ... {len(au.failures) - 12} more")
    charts, refs = audit_dashboard_chart(src)
    wbf = load_workbook(src)
    fmr = rows_by_label(wbf["Financial Model"])
    want = {f"'Financial Model'!$C${fmr['Passenger revenue']}:$N${fmr['Passenger revenue']}",
            f"'Financial Model'!$C${fmr['Route operating profit / (loss)']}:"
            f"$N${fmr['Route operating profit / (loss)']}",
            f"'Financial Model'!$C${fmr['Period']}:$N${fmr['Period']}"}
    chart_ok = bool(charts) and want <= set(refs)
    total_id += 1
    if not chart_ok:
        total_fail += 1
        detail.append(f"- Dashboard chart references {refs} do not cover {sorted(want)}")
    print(f"Dashboard chart: {len(charts)} chart(s); series/category refs {sorted(set(refs))} "
          f"-> {'linked to Financial Model' if chart_ok else 'NOT LINKED'}")
    lines.append("")
    lines.append(f"Dashboard chart references (read from the chart XML): {', '.join(sorted(set(refs)))}"
                 f" - {'linked to Financial Model revenue, route operating profit and period rows' if chart_ok else 'NOT LINKED'}.")
    print(f"\nTOTAL cells compared: {total_cmp}  identities tested: {total_id}  "
          f"mismatches: {total_fail}")
    if report:
        body = "\n".join(lines)
        if detail:
            body += "\n\nMismatches:\n" + "\n".join(detail)
        report.write_text(body + "\n")
    sys.exit(1 if total_fail else 0)


if __name__ == "__main__":
    main()
