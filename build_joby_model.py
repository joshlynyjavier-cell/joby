"""
Build the Joby Aviation illustrative air taxi route financial model (Excel).

ILLUSTRATIVE PROTOTYPE ONLY. Every number written as an input below is a
hypothetical assumption chosen for a case study. None of it is a Joby
Aviation forecast, guidance, or company-reported figure.

Usage:
    python build_joby_model.py [output_path]

Output defaults to Joby_Financial_Model.xlsx in this directory. Every
calculated cell is an Excel formula linked back to the Assumptions tab, and
the workbook is flagged to recalculate fully when opened. Run
validate_model.py afterwards to recalculate with LibreOffice and check
the results.

Layout conventions (kept identical across tabs so new tabs can plug in):
  * Assumptions: column F ("Active Case") is the single source every model
    formula reads. Columns C/D/E hold Base/Upside/Downside inputs.
  * Financial Model and the Scenario Comparison engine: months in C..N,
    annual total in O, annual basis label in P.
  * Key inputs and outputs are exposed as workbook-level defined names
    (inp_* / out_*), so a future Sensitivity Analysis or Executive Dashboard
    tab can reference them without hard-coding cell addresses.
"""

import datetime as dt
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

# --------------------------------------------------------------------------
# Style palette: navy / white / light blue
# --------------------------------------------------------------------------
NAVY = "1F3864"
MID_BLUE = "2F5597"
LIGHT_BLUE = "DDEBF7"
PALE_BLUE = "F2F7FC"
INPUT_FILL_HEX = "FFF2CC"  # pale yellow: marks cells the user may edit
WHITE = "FFFFFF"
GREY = "7F7F7F"

FONT_NAME = "Arial"
F_INPUT = Font(name=FONT_NAME, size=10, color="0000FF")  # editable input
F_CALC = Font(name=FONT_NAME, size=10, color="000000")  # formula
F_LINK = Font(name=FONT_NAME, size=10, color="008000")  # link to another tab
F_BODY = Font(name=FONT_NAME, size=10, color="000000")
F_BOLD = Font(name=FONT_NAME, size=10, bold=True, color="000000")
F_BOLD_CALC = Font(name=FONT_NAME, size=10, bold=True, color="000000")
F_NOTE = Font(name=FONT_NAME, size=9, italic=True, color=GREY)
F_TITLE = Font(name=FONT_NAME, size=16, bold=True, color=WHITE)
F_SUBTITLE = Font(name=FONT_NAME, size=10, italic=True, color=WHITE)
F_HEADER = Font(name=FONT_NAME, size=10, bold=True, color=WHITE)
F_SECTION = Font(name=FONT_NAME, size=10, bold=True, color=NAVY)
F_WARN = Font(name=FONT_NAME, size=10, bold=True, color="C00000")

FILL_NAVY = PatternFill("solid", fgColor=NAVY)
FILL_MID = PatternFill("solid", fgColor=MID_BLUE)
FILL_LIGHT = PatternFill("solid", fgColor=LIGHT_BLUE)
FILL_PALE = PatternFill("solid", fgColor=PALE_BLUE)
FILL_INPUT = PatternFill("solid", fgColor=INPUT_FILL_HEX)
FILL_PASS = PatternFill("solid", fgColor="C6EFCE")
FILL_FAIL = PatternFill("solid", fgColor="FFC7CE")

THIN = Side(style="thin", color="BFBFBF")
B_INPUT = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
B_TOP = Border(top=Side(style="thin", color=NAVY))
B_TOTAL = Border(top=Side(style="thin", color=NAVY), bottom=Side(style="double", color=NAVY))

A_LEFT = Alignment(horizontal="left", vertical="center")
A_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
A_CENTER = Alignment(horizontal="center", vertical="center")
A_RIGHT = Alignment(horizontal="right", vertical="center")

# Number formats (zeros render as "-", negatives in parentheses)
FMT_USD = '$#,##0;($#,##0);"-"'
FMT_USD2 = '$#,##0.00;($#,##0.00);"-"'
FMT_NUM = '#,##0;(#,##0);"-"'
FMT_NUM1 = '#,##0.0;(#,##0.0);"-"'
FMT_PCT = '0.0%;(0.0%);"-"'
FMT_PCT2 = '0.00%;(0.00%);"0.00%"'
FMT_IDX = '0.0000'
FMT_DATE = 'mmm-yy'
FMT_INT = '0'

DISCLAIMER = (
    "ILLUSTRATIVE PROTOTYPE - all inputs are hypothetical case-study assumptions, "
    "not Joby Aviation forecasts, guidance, or company-reported figures."
)

MONTHS = 12
FIRST_MONTH_COL = 3  # column C
MONTH_COLS = [get_column_letter(FIRST_MONTH_COL + i) for i in range(MONTHS)]  # C..N
FIRST_M, LAST_M = MONTH_COLS[0], MONTH_COLS[-1]
TOTAL_COL = get_column_letter(FIRST_MONTH_COL + MONTHS)  # O
BASIS_COL = get_column_letter(FIRST_MONTH_COL + MONTHS + 1)  # P

SCEN_COLS = {"Base": "C", "Upside": "D", "Downside": "E"}
ACTIVE_COL = "F"
AS = "Assumptions"
FM = "'Financial Model'"
SC = "'Scenario Comparison'"


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------
def style(cell, font=None, fill=None, fmt=None, align=None, border=None):
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if fmt is not None:
        cell.number_format = fmt
    if align is not None:
        cell.alignment = align
    if border is not None:
        cell.border = border
    return cell


def title_block(ws, title, subtitle, last_col):
    """Navy banner across rows 1-2 with a disclaimer line in row 3."""
    for row in (1, 2):
        for c in range(1, last_col + 1):
            ws.cell(row=row, column=c).fill = FILL_NAVY
    style(ws.cell(row=1, column=1, value=title), font=F_TITLE, align=A_LEFT)
    style(ws.cell(row=2, column=1, value=subtitle), font=F_SUBTITLE, align=A_LEFT)
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 18
    style(ws.cell(row=3, column=1, value=DISCLAIMER), font=F_WARN, align=A_LEFT)


def header_row(ws, row, values, last_col=None):
    last_col = last_col or len(values)
    for c in range(1, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = FILL_MID
        cell.font = F_HEADER
        cell.alignment = A_CENTER if c > 1 else A_LEFT
    for c, v in enumerate(values, start=1):
        ws.cell(row=row, column=c, value=v)
    ws.row_dimensions[row].height = 20


def section_row(ws, row, text, last_col):
    for c in range(1, last_col + 1):
        ws.cell(row=row, column=c).fill = FILL_LIGHT
    style(ws.cell(row=row, column=1, value=text), font=F_SECTION, align=A_LEFT)


def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


# --------------------------------------------------------------------------
# Assumption definitions (hypothetical, illustrative values)
# --------------------------------------------------------------------------
# key, label, unit, value(s), number format, source / notes
GENERAL_INPUTS = [
    ("route_name", "Route name", "text",
     "Route A: Downtown - Airport", "@",
     "Illustrative placeholder (e.g., downtown vertiport to regional airport, ~15-25 miles). Not an announced Joby route."),
    ("service_type", "Service / aircraft type", "text",
     "Piloted eVTOL air taxi", "@",
     "Illustrative 4-passenger piloted eVTOL configuration. Add a row per aircraft type when extending to a multi-aircraft fleet."),
    ("start_date", "Model start month", "date", dt.date(2027, 1, 1), FMT_DATE,
     "Illustrative timing used to label periods only. Not a Joby launch date."),
    ("capacity", "Aircraft passenger seats per flight", "seats / flight", 4, FMT_NUM,
     "Illustrative; consistent with a pilot + 4 passenger eVTOL layout. Verify against current company disclosures before external use."),
    ("flights_per_day", "Scheduled flights per operating day", "flights / day", 32, FMT_NUM,
     "Illustrative: e.g., 2 aircraft x 16 one-way flights over a ~14-hour operating window. One-way flights."),
    ("op_days", "Operating days per month", "days / month", 28, FMT_NUM,
     "Illustrative: allows ~2-3 non-operating days per month for weather, maintenance, holidays."),
]

# key, label, unit, (base, upside, downside), fmt, notes
SCENARIO_INPUTS = [
    ("completion", "Flight completion rate", "% of scheduled", (0.95, 0.97, 0.90), FMT_PCT,
     "Illustrative. Share of scheduled flights actually flown (weather, technical, ATC cancellations)."),
    ("fare", "Average fare per passenger (Month 1)", "$ / passenger", (175.0, 190.0, 160.0), FMT_USD2,
     "Illustrative premium-rideshare-style price point for a ~15-25 mile urban trip. Net of discounts, excl. taxes."),
    ("load_factor", "Passenger load factor (Month 1)", "% of seats", (0.60, 0.68, 0.52), FMT_PCT,
     "Illustrative. Passengers flown / available seats on completed flights in Month 1."),
    ("var_cost_flight", "Variable operating cost per completed flight (Month 1)", "$ / completed flight",
     (300.0, 280.0, 320.0), FMT_USD2,
     "Illustrative. Energy/charging, pilot cost per flight, maintenance reserves, landing & vertiport usage fees."),
    ("var_cost_pax", "Variable cost per passenger (Month 1)", "$ / passenger", (12.0, 10.0, 15.0), FMT_USD2,
     "Illustrative. Payment processing, booking/distribution, passenger handling and per-passenger fees."),
    ("fixed_monthly", "Monthly route-specific fixed costs (Month 1)", "$ / month",
     (90000.0, 85000.0, 100000.0), FMT_USD,
     "Illustrative. Route-dedicated vertiport lease, ground staff, local marketing. Excludes corporate overhead, R&D, certification, aircraft depreciation."),
]

GROWTH_INPUTS = [
    ("fare_growth", "Monthly fare growth rate", "% per month", (0.0, 0.0025, -0.005), FMT_PCT2,
     "Illustrative. Compounds monthly on the Month 1 fare. Negative = promotional pricing pressure."),
    ("demand_growth", "Monthly demand growth rate", "% per month", (0.015, 0.020, 0.005), FMT_PCT2,
     "Illustrative. Compounds monthly on the Month 1 load factor (modelled load factor capped at 100%)."),
    ("cost_growth", "Monthly operating cost growth rate", "% per month", (0.0, -0.0025, 0.005), FMT_PCT2,
     "Illustrative. Compounds monthly on variable cost per flight, variable cost per passenger and fixed costs. Negative = efficiency gains."),
]

SCENARIO_NARRATIVE = {
    "Base": "Steady launch: 95% completion, $175 fare, 60% load factor building ~1.5%/month, flat unit costs.",
    "Upside": "Strong adoption: higher fare and load factor, faster demand ramp, better reliability and modest cost efficiencies.",
    "Downside": "Soft launch: discounted fares, weaker demand, more weather/technical cancellations and cost inflation.",
}


# --------------------------------------------------------------------------
# Assumptions tab
# --------------------------------------------------------------------------
def build_assumptions(wb):
    ws = wb.create_sheet(AS)
    A = {}
    last_col = 7
    title_block(ws, "Assumptions - Illustrative Air Taxi Route",
                "Edit blue cells only. Column F (Active Case) feeds every model calculation.", last_col)

    # Scenario selector
    ws["A5"] = "Active scenario (select)"
    ws["A5"].font = F_BOLD
    style(ws["C5"], font=F_INPUT, fill=FILL_INPUT, align=A_CENTER, border=B_INPUT)
    ws["C5"] = "Base"
    ws["D5"] = "<- choose Base, Upside or Downside; drives the Financial Model tab"
    ws["D5"].font = F_NOTE
    ws["A6"] = "Active scenario index"
    ws["A6"].font = F_BODY
    ws["C6"] = "=MATCH($C$5,$C$9:$E$9,0)"
    style(ws["C6"], font=F_CALC, align=A_CENTER, fmt=FMT_INT)
    ws["A7"] = "Model integrity status"
    ws["A7"].font = F_BODY
    A["scenario"] = 5
    A["scenario_idx"] = 6
    A["status"] = 7  # filled after the Financial Model checks exist

    dv = DataValidation(type="list", formula1='"Base,Upside,Downside"', allow_blank=False,
                        showErrorMessage=True, errorTitle="Scenario",
                        error="Choose Base, Upside or Downside.")
    ws.add_data_validation(dv)
    dv.add("C5")

    # Legend (column G, rows 5-7 + row 4 label)
    ws["G4"] = "Legend"
    ws["G4"].font = F_BOLD
    style(ws["G5"], font=F_INPUT, fill=FILL_INPUT, border=B_INPUT)
    ws["G5"] = "Blue text / yellow fill = editable input (hypothetical, illustrative)"
    ws["G6"] = "Black text = formula (do not overwrite)"
    ws["G6"].font = F_CALC
    ws["G7"] = "Green text = link pulled from another tab"
    ws["G7"].font = F_LINK

    header_row(ws, 9, ["Parameter", "Unit", "Base", "Upside", "Downside",
                       "Active Case", "Source / Notes"], last_col)

    row = 10
    section_row(ws, row, "1. Route & Service Setup (applies to all scenarios)", last_col)
    row += 1
    for key, label, unit, value, fmt, note in GENERAL_INPUTS:
        ws.cell(row=row, column=1, value=label).font = F_BODY
        ws.cell(row=row, column=2, value=unit).font = F_BODY
        c = ws.cell(row=row, column=3, value=value)
        style(c, font=F_INPUT, fill=FILL_INPUT, fmt=fmt, border=B_INPUT,
              align=A_LEFT if unit == "text" else A_RIGHT)
        if unit == "text":
            # long text: let it spill across the scenario columns visually
            ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
            style(ws.cell(row=row, column=6, value=f"=$C{row}"), font=F_CALC, fmt=fmt,
                  align=A_LEFT)
        else:
            style(ws.cell(row=row, column=4, value="all cases"), font=F_NOTE, align=A_CENTER)
            style(ws.cell(row=row, column=5, value="all cases"), font=F_NOTE, align=A_CENTER)
            style(ws.cell(row=row, column=6, value=f"=$C{row}"), font=F_CALC, fmt=fmt,
                  align=A_RIGHT, fill=FILL_PALE)
        style(ws.cell(row=row, column=7, value=note), font=F_NOTE, align=A_WRAP)
        A[key] = row
        row += 1

    def scenario_block(title, items, row):
        section_row(ws, row, title, last_col)
        row += 1
        for key, label, unit, vals, fmt, note in items:
            ws.cell(row=row, column=1, value=label).font = F_BODY
            ws.cell(row=row, column=2, value=unit).font = F_BODY
            for col, v in zip("CDE", vals):
                style(ws[f"{col}{row}"], font=F_INPUT, fill=FILL_INPUT, fmt=fmt, border=B_INPUT,
                      align=A_RIGHT)
                ws[f"{col}{row}"] = v
            style(ws.cell(row=row, column=6,
                          value=f"=INDEX($C{row}:$E{row},$C${A['scenario_idx']})"),
                  font=F_BOLD_CALC, fmt=fmt, align=A_RIGHT, fill=FILL_PALE)
            style(ws.cell(row=row, column=7, value=note), font=F_NOTE, align=A_WRAP)
            A[key] = row
            row += 1
        return row

    row += 1
    row = scenario_block("2. Scenario Operating & Unit Economics Drivers", SCENARIO_INPUTS, row)
    row += 1
    row = scenario_block("3. Optional Monthly Growth Rates (compounding from Month 1)",
                         GROWTH_INPUTS, row)

    # Input validation (blocks impossible entries at the source)
    pct_dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1",
                            showErrorMessage=True, errorTitle="Percentage input",
                            error="Enter a rate between 0% and 100%.")
    nonneg_dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                               showErrorMessage=True, errorTitle="Non-negative input",
                               error="Enter a value of zero or more.")
    growth_dv = DataValidation(type="decimal", operator="between", formula1="-0.5",
                               formula2="0.5", showErrorMessage=True, errorTitle="Growth rate",
                               error="Enter a monthly rate between -50% and +50%.")
    for d in (pct_dv, nonneg_dv, growth_dv):
        ws.add_data_validation(d)
    for k in ("completion", "load_factor"):
        pct_dv.add(f"C{A[k]}:E{A[k]}")
    for k in ("fare", "var_cost_flight", "var_cost_pax", "fixed_monthly"):
        nonneg_dv.add(f"C{A[k]}:E{A[k]}")
    for k in ("capacity", "flights_per_day", "op_days"):
        nonneg_dv.add(f"C{A[k]}")
    for k in ("fare_growth", "demand_growth", "cost_growth"):
        growth_dv.add(f"C{A[k]}:E{A[k]}")

    row += 1
    section_row(ws, row, "4. Scenario Definitions", last_col)
    row += 1
    for name, text in SCENARIO_NARRATIVE.items():
        ws.cell(row=row, column=1, value=name).font = F_BOLD
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
        style(ws.cell(row=row, column=2, value=text), font=F_BODY, align=A_WRAP)
        row += 1

    row += 1
    section_row(ws, row, "5. Methodology Notes", last_col)
    row += 1
    notes = [
        "Only column F (Active Case) is referenced by the Financial Model; it picks Base/Upside/Downside via INDEX on the scenario index.",
        "Rates are stored as fractions (0.60 = 60%). Monthly growth compounds: value in month m = Month 1 value x (1 + rate)^(m - 1).",
        "Volumes are expected values (scheduled flights x completion rate), so completed flights and passengers may be fractional.",
        "Fixed costs here are route-specific only. Corporate overhead, R&D, certification and aircraft ownership costs are outside route operating profit.",
        "To add a route: copy columns C:F into a new route block (or add a Route ID dimension) and point a new Financial Model block at it.",
    ]
    for n in notes:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        style(ws.cell(row=row, column=1, value="- " + n), font=F_BODY, align=A_WRAP)
        ws.row_dimensions[row].height = 26
        row += 1

    widths = {"A": 50, "B": 20, "C": 15, "D": 15, "E": 15, "F": 16, "G": 90}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C10"
    ws.sheet_view.showGridLines = False
    return ws, A


def act(A, key, absolute=True):
    """Reference to the Active Case cell for an assumption."""
    return f"{AS}!$F${A[key]}"


def scen_ref(A, key, col):
    return f"{AS}!${col}${A[key]}"


# --------------------------------------------------------------------------
# Financial Model tab
# --------------------------------------------------------------------------
def build_financial_model(wb, A):
    ws = wb.create_sheet("Financial Model")
    R = {}
    last_col = 16
    title_block(ws, "Financial Model - Monthly Route Economics (12 months)",
                "Active scenario only. All values are formulas linked to the Assumptions tab.",
                last_col)
    ws["A4"] = (f'="Route: "&{act(A, "route_name")}&"   |   Aircraft: "&{act(A, "service_type")}'
                f'&"   |   Active scenario: "&{AS}!$C${A["scenario"]}')
    ws["A4"].font = Font(name=FONT_NAME, size=10, bold=True, color="008000")

    header_row(ws, 6, ["Line item", "Unit"] + [f"Month {i + 1}" for i in range(MONTHS)]
               + ["FY Total", "Annual basis"], last_col)
    # Month index and period date
    R["month"] = 7
    R["period"] = 8
    ws["A7"], ws["B7"] = "Month index", "#"
    ws["A8"], ws["B8"] = "Period", "month"
    for i, col in enumerate(MONTH_COLS):
        if i == 0:
            style(ws[f"{col}7"], font=F_CALC, fmt=FMT_INT, align=A_CENTER)
            ws[f"{col}7"] = 1
        else:
            ws[f"{col}7"] = f"={MONTH_COLS[i - 1]}7+1"
            style(ws[f"{col}7"], font=F_CALC, fmt=FMT_INT, align=A_CENTER)
        ws[f"{col}8"] = f"=EDATE({act(A, 'start_date')},{col}$7-1)"
        style(ws[f"{col}8"], font=F_LINK, fmt=FMT_DATE, align=A_CENTER)
    for r in (7, 8):
        for c in range(1, last_col + 1):
            ws.cell(row=r, column=c).fill = FILL_PALE
        ws.cell(row=r, column=1).font = F_BOLD

    row = 9

    def line(key, label, unit, month_formula, fmt, total=None, basis="", font=F_CALC,
             bold=False, border=None):
        """Write one model line. month_formula(col) -> formula string."""
        nonlocal row
        R[key] = row
        ws.cell(row=row, column=1, value=label).font = F_BOLD if bold else F_BODY
        ws.cell(row=row, column=2, value=unit).font = F_NOTE
        for col in MONTH_COLS:
            c = ws[f"{col}{row}"]
            c.value = month_formula(col)
            style(c, font=Font(name=FONT_NAME, size=10, bold=bold, color=font.color),
                  fmt=fmt, align=A_RIGHT, border=border)
        c = ws[f"{TOTAL_COL}{row}"]
        if total is not None:
            c.value = total
        # style the total cell even when its formula is written later
        style(c, font=Font(name=FONT_NAME, size=10, bold=True, color="000000"), fmt=fmt,
              align=A_RIGHT, fill=FILL_PALE, border=border)
        style(ws[f"{BASIS_COL}{row}"], font=F_NOTE, align=A_LEFT)
        ws[f"{BASIS_COL}{row}"] = basis
        row += 1

    def rng(key):
        return f"{FIRST_M}{R[key]}:{LAST_M}{R[key]}"

    def tot(key):
        return f"{TOTAL_COL}{R[key]}"

    def sum_total(key):
        return f"=SUM({rng(key)})"

    def section(text):
        nonlocal row
        row += 1
        section_row(ws, row, text, last_col)
        row += 1

    # ---- Driver build -----------------------------------------------------
    section("A. Driver Build (Active Case from Assumptions)")
    line("op_days", "Operating days", "days", lambda c: f"={act(A, 'op_days')}", FMT_NUM,
         None, "", font=F_LINK)
    ws[tot("op_days")] = sum_total("op_days")
    style(ws[tot("op_days")], font=F_BOLD, fmt=FMT_NUM, align=A_RIGHT)
    ws[f"{BASIS_COL}{R['op_days']}"] = "Sum"
    line("fpd", "Scheduled flights per operating day", "flights / day",
         lambda c: f"={act(A, 'flights_per_day')}", FMT_NUM, None, "", font=F_LINK)
    line("fare_idx", "Fare growth index", "index",
         lambda c: f"=(1+{act(A, 'fare_growth')})^({c}${R['month']}-1)", FMT_IDX)
    line("dem_idx", "Demand growth index", "index",
         lambda c: f"=(1+{act(A, 'demand_growth')})^({c}${R['month']}-1)", FMT_IDX)
    line("cost_idx", "Operating cost growth index", "index",
         lambda c: f"=(1+{act(A, 'cost_growth')})^({c}${R['month']}-1)", FMT_IDX)
    line("comp", "Flight completion rate", "% of scheduled",
         lambda c: f"={act(A, 'completion')}", FMT_PCT, None, "", font=F_LINK)
    line("lf", "Passenger load factor", "% of seats",
         lambda c: f"=MIN(1,MAX(0,{act(A, 'load_factor')}*{c}{R['dem_idx']}))", FMT_PCT)
    line("vcf", "Variable operating cost per completed flight", "$ / flight",
         lambda c: f"={act(A, 'var_cost_flight')}*{c}{R['cost_idx']}", FMT_USD2)
    line("vcp", "Variable cost per passenger", "$ / passenger",
         lambda c: f"={act(A, 'var_cost_pax')}*{c}{R['cost_idx']}", FMT_USD2)

    # ---- Operating metrics ---------------------------------------------------
    section("B. Operating Metrics")
    line("sched", "Scheduled flights", "flights",
         lambda c: f"={c}{R['op_days']}*{c}{R['fpd']}", FMT_NUM, None, "Sum")
    line("completed", "Completed flights", "flights",
         lambda c: f"={c}{R['sched']}*{c}{R['comp']}", FMT_NUM, None, "Sum")
    line("seats", "Available passenger seats", "seats",
         lambda c: f"={c}{R['completed']}*{act(A, 'capacity')}", FMT_NUM, None, "Sum")
    line("pax", "Passengers flown", "passengers",
         lambda c: f"={c}{R['seats']}*{c}{R['lf']}", FMT_NUM, None, "Sum", bold=True)
    for k in ("sched", "completed", "seats", "pax"):
        ws[tot(k)] = sum_total(k)

    # ---- Route P&L ----------------------------------------------------------
    section("C. Route P&L")
    line("fare", "Average fare", "$ / passenger",
         lambda c: f"={act(A, 'fare')}*{c}{R['fare_idx']}", FMT_USD2, None,
         "Weighted avg (revenue / passengers)")
    line("revenue", "Passenger revenue", "$",
         lambda c: f"={c}{R['pax']}*{c}{R['fare']}", FMT_USD, None, "Sum", bold=True)
    line("flight_cost", "Direct flight operating costs", "$",
         lambda c: f"={c}{R['completed']}*{c}{R['vcf']}", FMT_USD, None, "Sum")
    line("pax_cost", "Passenger-variable costs", "$",
         lambda c: f"={c}{R['pax']}*{c}{R['vcp']}", FMT_USD, None, "Sum")
    line("tvc", "Total variable costs", "$",
         lambda c: f"={c}{R['flight_cost']}+{c}{R['pax_cost']}", FMT_USD, None, "Sum",
         border=B_TOP)
    line("contrib", "Contribution profit (flight-level)", "$",
         lambda c: f"={c}{R['revenue']}-{c}{R['tvc']}", FMT_USD, None, "Sum", bold=True)
    line("cm_pct", "Contribution margin", "% of revenue",
         lambda c: f"=IF({c}{R['revenue']}=0,\"n.m.\",{c}{R['contrib']}/{c}{R['revenue']})", FMT_PCT,
         None, "Weighted (annual contribution / revenue)")
    line("contrib_flt", "Contribution profit per completed flight", "$ / flight",
         lambda c: f"=IF({c}{R['completed']}=0,\"n.m.\",{c}{R['contrib']}/{c}{R['completed']})",
         FMT_USD2, None, "Weighted (annual contribution / flights)")
    line("fixed", "Fixed route costs", "$",
         lambda c: f"={act(A, 'fixed_monthly')}*{c}{R['cost_idx']}", FMT_USD, None, "Sum")
    line("rop", "Route operating profit / (loss)", "$",
         lambda c: f"={c}{R['contrib']}-{c}{R['fixed']}", FMT_USD, None, "Sum", bold=True,
         border=B_TOTAL)
    line("rom", "Route operating margin", "% of revenue",
         lambda c: f"=IF({c}{R['revenue']}=0,\"n.m.\",{c}{R['rop']}/{c}{R['revenue']})", FMT_PCT,
         None, "Weighted (annual route op. profit / revenue)")
    line("cum_rop", "Cumulative route operating profit / (loss)", "$",
         lambda c: (f"={c}{R['rop']}" if c == FIRST_M else
                    f"={get_column_letter(FIRST_MONTH_COL + MONTH_COLS.index(c) - 1)}"
                    f"{R['cum_rop']}+{c}{R['rop']}"),
         FMT_USD, None, "End of period")

    for k in ("revenue", "flight_cost", "pax_cost", "tvc", "contrib", "fixed", "rop"):
        ws[tot(k)] = sum_total(k)
    ws[tot("fare")] = f"=IF({tot('pax')}=0,\"n.m.\",{tot('revenue')}/{tot('pax')})"
    ws[tot("cm_pct")] = f"=IF({tot('revenue')}=0,\"n.m.\",{tot('contrib')}/{tot('revenue')})"
    ws[tot("contrib_flt")] = f"=IF({tot('completed')}=0,\"n.m.\",{tot('contrib')}/{tot('completed')})"
    ws[tot("rom")] = f"=IF({tot('revenue')}=0,\"n.m.\",{tot('rop')}/{tot('revenue')})"
    ws[tot("cum_rop")] = f"={LAST_M}{R['cum_rop']}"

    # Weighted annual rates for the driver rows
    ws[tot("comp")] = f"=IF({tot('sched')}=0,\"n.m.\",{tot('completed')}/{tot('sched')})"
    ws[f"{BASIS_COL}{R['comp']}"] = "Weighted (completed / scheduled)"
    ws[tot("lf")] = f"=IF({tot('seats')}=0,\"n.m.\",{tot('pax')}/{tot('seats')})"
    ws[f"{BASIS_COL}{R['lf']}"] = "Weighted (passengers / seats)"
    ws[tot("vcf")] = f"=IF({tot('completed')}=0,\"n.m.\",{tot('flight_cost')}/{tot('completed')})"
    ws[f"{BASIS_COL}{R['vcf']}"] = "Weighted (flight costs / flights)"
    ws[tot("vcp")] = f"=IF({tot('pax')}=0,\"n.m.\",{tot('pax_cost')}/{tot('pax')})"
    ws[f"{BASIS_COL}{R['vcp']}"] = "Weighted (pax costs / passengers)"
    ws[tot("fpd")] = f"=IF({tot('op_days')}=0,\"n.m.\",{tot('sched')}/{tot('op_days')})"
    ws[f"{BASIS_COL}{R['fpd']}"] = "Weighted (flights / days)"
    for k, fmt in (("comp", FMT_PCT), ("lf", FMT_PCT), ("vcf", FMT_USD2), ("vcp", FMT_USD2),
                   ("fpd", FMT_NUM1)):
        style(ws[tot(k)], font=F_BOLD, fmt=fmt, align=A_RIGHT)

    # ---- Breakeven ----------------------------------------------------------
    section("D. Breakeven Analysis")
    line("net_pax_margin", "Net margin per passenger (fare - variable cost per passenger)",
         "$ / passenger", lambda c: f"={c}{R['fare']}-{c}{R['vcp']}", FMT_USD2, None,
         "Weighted ((revenue - pax costs) / passengers)")
    ws[tot("net_pax_margin")] = (f"=IF({tot('pax')}=0,\"n.m.\",({tot('revenue')}-{tot('pax_cost')})"
                                 f"/{tot('pax')})")
    line("max_margin", "Passenger margin if every seat were sold (100% load factor)", "$",
         lambda c: f"={c}{R['seats']}*{c}{R['net_pax_margin']}", FMT_USD, None, "Sum")
    ws[tot("max_margin")] = sum_total("max_margin")

    def be(c, include_fixed):
        costs = f"{c}{R['flight_cost']}" + (f"+{c}{R['fixed']}" if include_fixed else "")
        return f'=IF({c}{R["max_margin"]}>0,({costs})/{c}{R["max_margin"]},"n.m.")'

    line("be_flight", "Breakeven load factor - flight level (contribution = 0)", "% of seats",
         lambda c: be(c, False), FMT_PCT, None, "Annual: total flight costs / total max margin")
    ws[tot("be_flight")] = be(TOTAL_COL, False)
    line("be_route", "Breakeven load factor - route (after fixed route costs)", "% of seats",
         lambda c: be(c, True), FMT_PCT, None,
         "Annual: (flight + fixed costs) / total max margin", bold=True)
    ws[tot("be_route")] = be(TOTAL_COL, True)
    line("lf_headroom", "Load factor headroom vs. route breakeven", "pct. points",
         lambda c: f'=IF(ISNUMBER({c}{R["be_route"]}),{c}{R["lf"]}-{c}{R["be_route"]},"n.m.")',
         FMT_PCT, None, "Positive = above breakeven")
    ws[tot("lf_headroom")] = (f'=IF(ISNUMBER({tot("be_route")}),{tot("lf")}-{tot("be_route")},'
                              f'"n.m.")')
    for k in ("net_pax_margin", "be_flight", "be_route", "lf_headroom"):
        fmt = FMT_USD2 if k == "net_pax_margin" else FMT_PCT
        style(ws[tot(k)], font=F_BOLD, fmt=fmt, align=A_RIGHT, fill=FILL_PALE)
    row += 0
    note_row = row
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=last_col)
    style(ws.cell(row=note_row, column=1,
                  value=("Breakeven load factor = (direct flight costs [+ fixed route costs]) / "
                         "(available seats x (fare - variable cost per passenger)), holding fare, "
                         "completion and unit costs at that period's values. Values above 100% "
                         "mean breakeven is not reachable on current capacity. 'n.m.' = not "
                         "meaningful (fare does not cover variable cost per passenger).")),
          font=F_NOTE, align=A_WRAP)
    ws.row_dimensions[note_row].height = 28
    row += 1

    # Conditional formats: highlight unreachable breakeven and losses
    for k in ("be_flight", "be_route"):
        ws.conditional_formatting.add(
            f"{FIRST_M}{R[k]}:{TOTAL_COL}{R[k]}",
            FormulaRule(formula=[f'AND(ISNUMBER({FIRST_M}{R[k]}),{FIRST_M}{R[k]}>1)'],
                        font=Font(color="C00000", bold=True)))
    ws.conditional_formatting.add(
        f"{FIRST_M}{R['rop']}:{TOTAL_COL}{R['rop']}",
        CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))

    return ws, R, row, last_col


# --------------------------------------------------------------------------
# Scenario Comparison tab (computes all three scenarios side by side)
# --------------------------------------------------------------------------
def build_scenarios(wb, A):
    ws = wb.create_sheet("Scenario Comparison")
    last_col = 16
    title_block(ws, "Scenario Comparison - Base / Upside / Downside",
                "All three cases are computed simultaneously in the engine below, independent "
                "of the active-scenario selector.", last_col)
    S = {}

    # ---- Engine (built first so the summary can reference its rows) ---------
    engine_start = 50
    E = {s: {} for s in SCEN_COLS}
    row = engine_start
    section_row(ws, row, "Scenario Calculation Engine (monthly; same logic as the Financial "
                          "Model tab)", last_col)
    row += 1

    def eline(s, key, label, unit, fn, fmt, total=None):
        nonlocal row
        E[s][key] = row
        ws.cell(row=row, column=1, value=label).font = F_BODY
        ws.cell(row=row, column=2, value=unit).font = F_NOTE
        for c in MONTH_COLS:
            style(ws[f"{c}{row}"], font=F_CALC, fmt=fmt, align=A_RIGHT)
            ws[f"{c}{row}"] = fn(c)
        if total == "SUM":
            total = f"=SUM({FIRST_M}{row}:{LAST_M}{row})"
        if total:
            style(ws[f"{TOTAL_COL}{row}"], font=F_BOLD, fmt=fmt, align=A_RIGHT, fill=FILL_PALE)
            ws[f"{TOTAL_COL}{row}"] = total
        else:
            ws[f"{TOTAL_COL}{row}"].fill = FILL_PALE
        row += 1

    for s, sc in SCEN_COLS.items():
        row += 1
        header_row(ws, row, [f"{s} case", "Unit"] + [f"Month {i + 1}" for i in range(MONTHS)]
                   + ["FY Total", ""], last_col)
        row += 1
        r0 = row
        E[s]["month"] = row
        ws.cell(row=row, column=1, value="Month index").font = F_BODY
        for i, c in enumerate(MONTH_COLS):
            ws[f"{c}{row}"] = 1 if i == 0 else f"={MONTH_COLS[i - 1]}{row}+1"
            style(ws[f"{c}{row}"], font=F_CALC, fmt=FMT_INT, align=A_CENTER)
        row += 1
        m = lambda c: f"{c}${r0}"  # noqa: E731
        g = lambda k: scen_ref(A, k, sc)  # noqa: E731
        e = E[s]

        def r(k, c):
            return f"{c}{e[k]}"

        def sm(k):
            return f"=SUM({FIRST_M}{e[k]}:{LAST_M}{e[k]})"

        eline(s, "fare", "Average fare", "$ / passenger",
              lambda c: f"={g('fare')}*(1+{g('fare_growth')})^({m(c)}-1)", FMT_USD2)
        eline(s, "lf", "Passenger load factor", "% of seats",
              lambda c: f"=MIN(1,MAX(0,{g('load_factor')}*(1+{g('demand_growth')})^({m(c)}-1)))",
              FMT_PCT)
        eline(s, "cost_idx", "Operating cost growth index", "index",
              lambda c: f"=(1+{g('cost_growth')})^({m(c)}-1)", FMT_IDX)
        eline(s, "sched", "Scheduled flights", "flights",
              lambda c: f"={act(A, 'op_days')}*{act(A, 'flights_per_day')}", FMT_NUM, None)
        ws[f"{TOTAL_COL}{e['sched']}"] = sm("sched")
        style(ws[f"{TOTAL_COL}{e['sched']}"], font=F_BOLD, fmt=FMT_NUM, align=A_RIGHT,
              fill=FILL_PALE)
        eline(s, "completed", "Completed flights", "flights",
              lambda c: f"={r('sched', c)}*{g('completion')}", FMT_NUM, "SUM")
        eline(s, "seats", "Available passenger seats", "seats",
              lambda c: f"={r('completed', c)}*{act(A, 'capacity')}", FMT_NUM, "SUM")
        eline(s, "pax", "Passengers flown", "passengers",
              lambda c: f"={r('seats', c)}*{r('lf', c)}", FMT_NUM, "SUM")
        eline(s, "revenue", "Passenger revenue", "$",
              lambda c: f"={r('pax', c)}*{r('fare', c)}", FMT_USD, "SUM")
        eline(s, "flight_cost", "Direct flight operating costs", "$",
              lambda c: f"={r('completed', c)}*{g('var_cost_flight')}*{r('cost_idx', c)}",
              FMT_USD, "SUM")
        eline(s, "pax_cost", "Passenger-variable costs", "$",
              lambda c: f"={r('pax', c)}*{g('var_cost_pax')}*{r('cost_idx', c)}",
              FMT_USD, "SUM")
        eline(s, "tvc", "Total variable costs", "$",
              lambda c: f"={r('flight_cost', c)}+{r('pax_cost', c)}", FMT_USD, "SUM")
        eline(s, "contrib", "Contribution profit", "$",
              lambda c: f"={r('revenue', c)}-{r('tvc', c)}", FMT_USD, "SUM")
        eline(s, "fixed", "Fixed route costs", "$",
              lambda c: f"={g('fixed_monthly')}*{r('cost_idx', c)}", FMT_USD, "SUM")
        eline(s, "rop", "Route operating profit / (loss)", "$",
              lambda c: f"={r('contrib', c)}-{r('fixed', c)}", FMT_USD, "SUM")
        eline(s, "max_margin", "Passenger margin at 100% load factor", "$",
              lambda c: f"={r('seats', c)}*({r('fare', c)}-{g('var_cost_pax')}*{r('cost_idx', c)})",
              FMT_USD, "SUM")
        ws.conditional_formatting.add(
            f"{FIRST_M}{e['rop']}:{TOTAL_COL}{e['rop']}",
            CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))

    # ---- Summary table ------------------------------------------------------
    header_row(ws, 5, ["Metric", "Unit", "Base", "Upside", "Downside", "Upside vs Base",
                       "Downside vs Base", "Notes"], last_col)
    row = 6
    section_row(ws, row, "Key Scenario Assumptions (linked from Assumptions tab)", last_col)
    row += 1
    assumption_rows = [
        ("completion", "Flight completion rate", "%", FMT_PCT),
        ("fare", "Average fare (Month 1)", "$ / passenger", FMT_USD2),
        ("load_factor", "Load factor (Month 1)", "% of seats", FMT_PCT),
        ("var_cost_flight", "Variable cost per completed flight (Month 1)", "$ / flight", FMT_USD2),
        ("var_cost_pax", "Variable cost per passenger (Month 1)", "$ / passenger", FMT_USD2),
        ("fixed_monthly", "Fixed route costs (Month 1)", "$ / month", FMT_USD),
        ("fare_growth", "Monthly fare growth", "% / month", FMT_PCT2),
        ("demand_growth", "Monthly demand growth", "% / month", FMT_PCT2),
        ("cost_growth", "Monthly operating cost growth", "% / month", FMT_PCT2),
    ]
    for key, label, unit, fmt in assumption_rows:
        ws.cell(row=row, column=1, value=label).font = F_BODY
        ws.cell(row=row, column=2, value=unit).font = F_NOTE
        for s, sc in SCEN_COLS.items():
            style(ws[f"{sc}{row}"], font=F_LINK, fmt=fmt, align=A_RIGHT)
            ws[f"{sc}{row}"] = f"={scen_ref(A, key, sc)}"
        for vc, sc in (("F", "D"), ("G", "E")):
            style(ws[f"{vc}{row}"], font=F_CALC, fmt=fmt, align=A_RIGHT)
            ws[f"{vc}{row}"] = f"={sc}{row}-C{row}"
        row += 1

    row += 1
    section_row(ws, row, "Annual Results (12 months)", last_col)
    row += 1

    def tref(s, k):
        return f"{TOTAL_COL}{E[s][k]}"

    results = [
        ("sched", "Scheduled flights", "flights", FMT_NUM, lambda s: f"={tref(s, 'sched')}", ""),
        ("completed", "Completed flights", "flights", FMT_NUM,
         lambda s: f"={tref(s, 'completed')}", ""),
        ("pax", "Passengers flown", "passengers", FMT_NUM, lambda s: f"={tref(s, 'pax')}", ""),
        ("lf", "Average load factor (weighted)", "% of seats", FMT_PCT,
         lambda s: f"=IF({tref(s, 'seats')}=0,\"n.m.\",{tref(s, 'pax')}/{tref(s, 'seats')})",
         "Passengers / available seats"),
        ("fare", "Average fare (weighted)", "$ / passenger", FMT_USD2,
         lambda s: f"=IF({tref(s, 'pax')}=0,\"n.m.\",{tref(s, 'revenue')}/{tref(s, 'pax')})",
         "Revenue / passengers"),
        ("revenue", "Annual passenger revenue", "$", FMT_USD,
         lambda s: f"={tref(s, 'revenue')}", ""),
        ("tvc", "Total variable costs", "$", FMT_USD, lambda s: f"={tref(s, 'tvc')}",
         "Direct flight + passenger-variable costs"),
        ("contrib", "Contribution profit", "$", FMT_USD, lambda s: f"={tref(s, 'contrib')}",
         "Flight-level, before fixed route costs"),
        ("cm_pct", "Contribution margin", "% of revenue", FMT_PCT,
         lambda s: f"=IF({tref(s, 'revenue')}=0,\"n.m.\",{tref(s, 'contrib')}/{tref(s, 'revenue')})", ""),
        ("contrib_flt", "Contribution profit per completed flight", "$ / flight", FMT_USD2,
         lambda s: f"=IF({tref(s, 'completed')}=0,\"n.m.\",{tref(s, 'contrib')}/{tref(s, 'completed')})",
         ""),
        ("fixed", "Fixed route costs", "$", FMT_USD, lambda s: f"={tref(s, 'fixed')}", ""),
        ("rop", "Route operating profit / (loss)", "$", FMT_USD,
         lambda s: f"={tref(s, 'rop')}", "After route fixed costs; not EBITDA"),
        ("rom", "Route operating margin", "% of revenue", FMT_PCT,
         lambda s: f"=IF({tref(s, 'revenue')}=0,\"n.m.\",{tref(s, 'rop')}/{tref(s, 'revenue')})", ""),
        ("be_flight", "Breakeven load factor - flight level", "% of seats", FMT_PCT,
         lambda s: (f'=IF({tref(s, "max_margin")}>0,{tref(s, "flight_cost")}/'
                    f'{tref(s, "max_margin")},"n.m.")'),
         "Load factor at which contribution = 0"),
        ("be_route", "Breakeven load factor - route", "% of seats", FMT_PCT,
         lambda s: (f'=IF({tref(s, "max_margin")}>0,({tref(s, "flight_cost")}+{tref(s, "fixed")})/'
                    f'{tref(s, "max_margin")},"n.m.")'),
         "Load factor at which route operating profit = 0"),
        ("profit_months", "Months with positive route operating profit", "months", FMT_INT,
         lambda s: f'=COUNTIF({FIRST_M}{E[s]["rop"]}:{LAST_M}{E[s]["rop"]},">0")', ""),
    ]
    key_bold = {"revenue", "contrib", "rop", "contrib_flt", "be_route"}
    for key, label, unit, fmt, fn, note in results:
        S[key] = row
        bold = key in key_bold
        ws.cell(row=row, column=1, value=label).font = F_BOLD if bold else F_BODY
        ws.cell(row=row, column=2, value=unit).font = F_NOTE
        for s, sc in SCEN_COLS.items():
            style(ws[f"{sc}{row}"], font=F_BOLD_CALC if bold else F_CALC, fmt=fmt, align=A_RIGHT)
            ws[f"{sc}{row}"] = fn(s)
        for vc, sc in (("F", "D"), ("G", "E")):
            style(ws[f"{vc}{row}"], font=F_CALC, fmt=fmt, align=A_RIGHT)
            ws[f"{vc}{row}"] = (f'=IF(AND(ISNUMBER({sc}{row}),ISNUMBER(C{row})),{sc}{row}-C{row},'
                                f'"n.m.")')
        style(ws[f"H{row}"], font=F_NOTE, align=A_LEFT)
        ws[f"H{row}"] = note
        row += 1
    ws.conditional_formatting.add(
        f"C{S['rop']}:G{S['rop']}",
        CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))
    ws.conditional_formatting.add(
        f"C{S['contrib']}:G{S['contrib_flt']}",
        CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))

    row += 1
    section_row(ws, row, "Scenario Definitions", last_col)
    row += 1
    for name, text in SCENARIO_NARRATIVE.items():
        ws.cell(row=row, column=1, value=name).font = F_BOLD
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_col)
        style(ws.cell(row=row, column=2, value=text), font=F_BODY, align=A_LEFT)
        row += 1
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    style(ws.cell(row=row, column=1,
                  value=("Variance columns show Upside/Downside minus Base (percentage rows are "
                         "differences in percentage points). Scenario values are illustrative.")),
          font=F_NOTE, align=A_WRAP)
    assert row < engine_start, "summary overlaps the engine block"

    widths = {"A": 48, "B": 15}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for c in MONTH_COLS + [TOTAL_COL]:
        ws.column_dimensions[c].width = 14
    ws.column_dimensions["F"].width = 17  # variance headers in the summary
    ws.column_dimensions["G"].width = 17
    ws.column_dimensions["H"].width = 14  # month column in the engine; notes may overflow
    ws.column_dimensions[BASIS_COL].width = 4
    ws.freeze_panes = "C6"
    ws.sheet_view.showGridLines = False
    return ws, S, E


# --------------------------------------------------------------------------
# Sensitivity Analysis tab: fare x load factor -> contribution per completed flight
# --------------------------------------------------------------------------
SA = "'Sensitivity Analysis'"
SA_FARE_STEPS_EACH_SIDE = 4  # 9 fare columns centred on the Base fare
SA_LF_ROWS = 15  # 30% .. 100% in 5-point steps by default
FILL_AMBER = PatternFill("solid", fgColor="FFEB9C")


def build_sensitivity(wb, A, E):
    ws = wb.create_sheet("Sensitivity Analysis")
    last_col = 11  # K
    title_block(ws, "Sensitivity Analysis - Contribution per Completed Flight",
                "Base scenario, Month 1 unit economics. Average fare (columns) x passenger load "
                "factor (rows). Linked to the Assumptions tab.", last_col)
    X = {}
    base = lambda k: f"{AS}!$C${A[k]}"  # noqa: E731  Base column (general inputs live in C too)
    eb = E["Base"]
    eng = lambda k: f"{SC}!${FIRST_M}${eb[k]}"  # noqa: E731  Base engine, Month 1

    def put(row, label, formula, fmt, unit, note, font=F_LINK, key=None, bold=False):
        style(ws.cell(row=row, column=1, value=label), font=F_BOLD if bold else F_BODY)
        style(ws.cell(row=row, column=2, value=formula), font=font, fmt=fmt, align=A_RIGHT)
        if font is F_INPUT:
            style(ws.cell(row=row, column=2), fill=FILL_INPUT, border=B_INPUT)
        style(ws.cell(row=row, column=3, value=unit), font=F_NOTE)
        style(ws.cell(row=row, column=4, value=note), font=F_NOTE, align=A_LEFT)
        if key:
            X[key] = row

    def b(key):
        return f"$B${X[key]}"

    # ---- Linked Base inputs ---------------------------------------------------
    section_row(ws, 5, "1. Base Scenario Inputs (Month 1, linked from Assumptions - edit there)",
                last_col)
    put(6, "Aircraft passenger seats per flight", f"={base('capacity')}", FMT_NUM,
        "seats / flight", "Assumptions, Route setup (all scenarios).", key="cap")
    put(7, "Base average fare", f"={base('fare')}", FMT_USD2, "$ / passenger",
        "Assumptions, Base column. Centre column of the matrix.", key="fare")
    put(8, "Base passenger load factor", f"={base('load_factor')}", FMT_PCT, "% of seats",
        "Assumptions, Base column (Month 1, before demand growth).", key="lf")
    put(9, "Direct flight operating cost per completed flight", f"={base('var_cost_flight')}",
        FMT_USD2, "$ / flight", "Assumptions, Base column.", key="vcf")
    put(10, "Variable cost per passenger", f"={base('var_cost_pax')}", FMT_USD2,
        "$ / passenger", "Assumptions, Base column.", key="vcp")
    put(11, "Fixed route costs per month", f"={base('fixed_monthly')}", FMT_USD, "$ / month",
        "Assumptions, Base column. Used only for route-level breakeven.", key="fixed")
    put(12, "Completed flights per month",
        f"={base('op_days')}*{base('flights_per_day')}*{base('completion')}", FMT_NUM1,
        "flights / month", "Operating days x flights per day x Base completion rate.",
        font=F_CALC, key="completed")

    # ---- Matrix settings ------------------------------------------------------
    section_row(ws, 14, "2. Matrix Settings (editable, illustrative)", last_col)
    put(15, "Fare step between columns", 15, FMT_USD2, "$ / passenger",
        f"Columns run {SA_FARE_STEPS_EACH_SIDE} steps below to {SA_FARE_STEPS_EACH_SIDE} steps "
        "above the Base fare (floored at $0).", font=F_INPUT, key="fare_step")
    put(16, "Load factor - first row", 0.30, FMT_PCT, "% of seats",
        "Top row of the matrix.", font=F_INPUT, key="lf_start")
    put(17, "Load factor step between rows", 0.05, FMT_PCT, "pct. points",
        f"{SA_LF_ROWS} rows; keep the last row at or below 100% (checked below).",
        font=F_INPUT, key="lf_step")
    put(18, "Breakeven band (+/-)", 10, FMT_USD2, "$ / flight",
        "Combinations within this band of $0 are shaded amber as 'breakeven'.",
        font=F_INPUT, key="band")
    put(19, "Illustrative direct flight cost change", 25, FMT_USD2, "$ / flight",
        "Used in insight 3 only.", font=F_INPUT, key="cost_shock")
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1",
                        showErrorMessage=True, error="Enter a rate between 0% and 100%.")
    nn = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                        showErrorMessage=True, error="Enter a value of zero or more.")
    ws.add_data_validation(dv)
    ws.add_data_validation(nn)
    dv.add(f"B{X['lf_start']}:B{X['lf_step']}")
    nn.add(f"B{X['fare_step']}")
    nn.add(f"B{X['band']}:B{X['cost_shock']}")

    # ---- Breakeven ------------------------------------------------------------
    section_row(ws, 21, "3. Breakeven Load Factor (Base, Month 1)", last_col)
    margin = f"({b('cap')}*({b('fare')}-{b('vcp')}))"  # $ per flight per 100% LF
    put(22, "FLIGHT-LEVEL breakeven load factor (contribution per flight = $0)",
        f'=IF({b("fare")}>{b("vcp")},{b("vcf")}/{margin},"n.m.")', FMT_PCT, "% of seats",
        "Formula: Direct flight cost / (Seats x (Fare - Variable cost per passenger)). Covers "
        "flight-level variable costs only.", font=F_CALC, key="be_flight", bold=True)
    put(23, "Fixed route costs per completed flight",
        f'=IF({b("completed")}>0,{b("fixed")}/{b("completed")},"n.m.")', FMT_USD2,
        "$ / flight", "Formula: Monthly fixed route costs / completed flights per month.",
        font=F_CALC, key="fixed_pf")
    put(24, "ROUTE-LEVEL breakeven load factor (after fixed route costs)",
        f'=IF(AND({b("fare")}>{b("vcp")},ISNUMBER({b("fixed_pf")})),'
        f'({b("vcf")}+{b("fixed_pf")})/{margin},"n.m.")', FMT_PCT, "% of seats",
        "Formula: (Direct flight cost + Fixed route cost per flight) / (Seats x (Fare - Variable "
        "cost per passenger)). Route operating profit = $0.", font=F_CALC, key="be_route",
        bold=True)
    put(25, "Gap: route-level minus flight-level breakeven",
        f'=IF(AND(ISNUMBER({b("be_route")}),ISNUMBER({b("be_flight")})),'
        f'{b("be_route")}-{b("be_flight")},"n.m.")', FMT_PCT, "pct. points",
        "Extra load factor needed to cover fixed route costs.", font=F_CALC, key="be_gap")
    put(26, "Base load factor headroom vs. flight-level breakeven",
        f'=IF(ISNUMBER({b("be_flight")}),{b("lf")}-{b("be_flight")},"n.m.")', FMT_PCT,
        "pct. points", "Positive = Base case covers its direct flight costs.", font=F_CALC,
        key="hr_flight")
    put(27, "Base load factor headroom vs. route-level breakeven",
        f'=IF(ISNUMBER({b("be_route")}),{b("lf")}-{b("be_route")},"n.m.")', FMT_PCT,
        "pct. points", "Positive = Base case also covers fixed route costs (Month 1).",
        font=F_CALC, key="hr_route")
    for k in ("be_flight", "be_route", "be_gap", "hr_flight", "hr_route"):
        ws.cell(row=X[k], column=2).fill = FILL_PALE

    # ---- Matrix -----------------------------------------------------------------
    section_row(ws, 29, "4. Contribution Profit per Completed Flight ($) - Average Fare "
                        "(columns) x Load Factor (rows)", last_col)
    hdr = 30
    first_r, last_r = hdr + 1, hdr + SA_LF_ROWS
    fcols = [get_column_letter(3 + i) for i in range(2 * SA_FARE_STEPS_EACH_SIDE + 1)]  # C..K
    for c in range(1, last_col + 1):
        ws.cell(row=hdr, column=c).fill = FILL_MID
    style(ws.cell(row=hdr, column=1, value="Formula: Seats x LF x (Fare - Var. cost/pax) - "
                                           "Direct flight cost"),
          font=Font(name=FONT_NAME, size=9, italic=True, color=WHITE), align=A_LEFT)
    style(ws.cell(row=hdr, column=2, value="LF \\ Fare"), font=F_HEADER, align=A_CENTER)
    for i, col in enumerate(fcols):
        if i == 0:
            f = f"=MAX(0,ROUND({b('fare')}-{SA_FARE_STEPS_EACH_SIDE}*{b('fare_step')},2))"
        else:
            f = f"=MAX(0,ROUND({fcols[i - 1]}{hdr}+{b('fare_step')},2))"
        style(ws[f"{col}{hdr}"], font=F_HEADER, fmt=FMT_USD2, align=A_CENTER)
        ws[f"{col}{hdr}"] = f
    for j in range(SA_LF_ROWS):
        r = first_r + j
        f = (f"=ROUND({b('lf_start')},4)" if j == 0 else f"=ROUND(B{r - 1}+{b('lf_step')},4)")
        style(ws[f"B{r}"], font=F_BOLD, fmt=FMT_PCT, align=A_CENTER, fill=FILL_LIGHT)
        ws[f"B{r}"] = f
        for col in fcols:
            style(ws[f"{col}{r}"], font=F_CALC, fmt=FMT_USD, align=A_RIGHT, border=B_INPUT)
            ws[f"{col}{r}"] = (f"={b('cap')}*$B{r}*({col}${hdr}-{b('vcp')})-{b('vcf')}")
    grid = f"$C${first_r}:$K${last_r}"
    lf_axis = f"$B${first_r}:$B${last_r}"
    fare_axis = f"$C${hdr}:$K${hdr}"
    tl = f"C{first_r}"
    ws.conditional_formatting.add(
        grid, FormulaRule(formula=[f"ABS({tl})<={b('band')}"], fill=FILL_AMBER, stopIfTrue=True))
    ws.conditional_formatting.add(
        grid, FormulaRule(formula=[f"{tl}<-{b('band')}"], fill=FILL_FAIL,
                          font=Font(color="9C0006")))
    ws.conditional_formatting.add(
        grid, FormulaRule(formula=[f"{tl}>{b('band')}"], fill=FILL_PASS,
                          font=Font(color="006100")))
    ws.conditional_formatting.add(
        grid, FormulaRule(formula=[f"AND(C${hdr}={b('fare')},$B{first_r}={b('lf')})"],
                          font=Font(bold=True, underline="single"),
                          border=Border(left=Side(style="medium", color=NAVY),
                                        right=Side(style="medium", color=NAVY),
                                        top=Side(style="medium", color=NAVY),
                                        bottom=Side(style="medium", color=NAVY))))
    X.update(hdr=hdr, first_r=first_r, last_r=last_r)
    # Base-case callout beside the matrix (computed directly, not looked up)
    style(ws.cell(row=first_r, column=1, value="Base case (Month 1):"), font=F_BOLD, align=A_RIGHT)
    style(ws.cell(row=first_r + 1, column=1,
                  value=(f'=TEXT({b("fare")},"$0.00")&" fare x "&TEXT({b("lf")},"0.0%")&'
                         f'" load factor = "&TEXT({b("cap")}*{b("lf")}*({b("fare")}-{b("vcp")})'
                         f'-{b("vcf")},"$#,##0.00")&" per completed flight"')),
          font=F_CALC, align=A_RIGHT)
    style(ws.cell(row=first_r + 2, column=1, value="(bold, boxed cell in the grid when on the axes)"),
          font=F_NOTE, align=A_RIGHT)

    # Breakeven by fare column
    r = last_r + 1
    X["be_flight_row"], X["be_route_row"] = r, r + 1
    for rr, label, extra in ((r, "Flight-level breakeven LF at this fare", ""),
                             (r + 1, "Route-level breakeven LF at this fare (after fixed costs)",
                              f"+{b('fixed_pf')}")):
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=2)
        style(ws.cell(row=rr, column=1, value=label), font=F_BOLD, align=A_RIGHT, fill=FILL_PALE)
        for col in fcols:
            cond = f"{col}${hdr}>{b('vcp')}"
            if extra:
                cond = f"AND({cond},ISNUMBER({b('fixed_pf')}))"
            style(ws[f"{col}{rr}"], font=F_BOLD_CALC, fmt=FMT_PCT, align=A_RIGHT, fill=FILL_PALE,
                  border=B_TOP)
            ws[f"{col}{rr}"] = (f'=IF({cond},({b("vcf")}{extra})/({b("cap")}*({col}${hdr}-'
                                f'{b("vcp")})),"n.m.")')
        ws.conditional_formatting.add(
            f"C{rr}:K{rr}", FormulaRule(formula=[f"AND(ISNUMBER(C{rr}),C{rr}>1)"],
                                        font=Font(color="C00000", bold=True)))
    r += 2
    style(ws.cell(row=r, column=1,
                  value="Bold boxed cell = Base fare x Base load factor (when on the grid). "
                        "Breakeven LF above 100% (red) cannot be reached at that fare."),
          font=F_NOTE)

    # Legend and classification counts
    r += 2
    X["legend"] = r
    band = b("band")
    legend = [
        ("Loss-making (below -band)", FILL_FAIL, f'=COUNTIF({grid},"<"&-{band})'),
        ("Breakeven (within +/- band)", FILL_AMBER,
         f'=COUNTIFS({grid},">="&-{band},{grid},"<="&{band})'),
        ("Profitable (above +band)", FILL_PASS, f'=COUNTIF({grid},">"&{band})'),
    ]
    style(ws.cell(row=r, column=1, value="Legend (fare / load factor combinations)"),
          font=F_BOLD)
    style(ws.cell(row=r, column=2, value="Count"), font=F_BOLD, align=A_CENTER)
    for i, (label, fill, f) in enumerate(legend, start=1):
        style(ws.cell(row=r + i, column=1, value=label), font=F_BODY, fill=fill)
        style(ws.cell(row=r + i, column=2, value=f), font=F_CALC, fmt=FMT_INT, align=A_CENTER)
    X["cnt_first"], X["cnt_last"] = r + 1, r + 3
    r += 5

    # ---- Insights ---------------------------------------------------------------
    section_row(ws, r, "5. Business Insights (illustrative; text updates with the inputs)",
                last_col)
    r += 1
    cap, fare, lf, vcf, vcp = b("cap"), b("fare"), b("lf"), b("vcf"), b("vcp")
    low_be = f"$C${X['be_flight_row']}"
    high_be = f"$K${X['be_flight_row']}"
    insights = [
        ('="1. Pricing: at the Base "&TEXT(' + lf + ',"0%")&" load factor, every "&TEXT('
         + b("fare_step") + ',"$0")&" change in average fare moves contribution per completed '
         'flight by "&TEXT(' + cap + "*" + lf + "*" + b("fare_step") + ',"$#,##0")&". Across the '
         'matrix the flight-level breakeven load factor falls from "&TEXT(' + low_be
         + ',"0.0%")&" at "&TEXT($C$' + str(hdr) + ',"$0")&" to "&TEXT(' + high_be
         + ',"0.0%")&" at "&TEXT($K$' + str(hdr) + ',"$0")&", so pricing power directly widens '
         'the margin of safety against weak demand."'),
        ('="2. Demand: each "&TEXT(' + b("lf_step") + '*100,"0")&"-point change in load factor '
         'moves contribution per flight by "&TEXT(' + cap + "*" + b("lf_step") + "*(" + fare
         + "-" + vcp + '),"$#,##0")&". The Base route covers its direct flight costs above a "'
         '&TEXT(' + b("be_flight") + ',"0.0%")&" load factor but needs "&TEXT('
         + b("be_route") + ',"0.0%")&" to also cover fixed route costs - a "&TEXT('
         + b("be_gap") + '*100,"0.0")&"-point gap that only more completed flights, a higher '
         'fare or lower fixed costs can close."'),
        ('="3. Cost: a "&TEXT(' + b("cost_shock") + ',"$0")&" increase in direct flight cost '
         'per completed flight cuts contribution per flight by the same "&TEXT('
         + b("cost_shock") + ',"$0")&" at every fare and load factor and raises the flight-level '
         'breakeven load factor by "&TEXT(' + b("cost_shock") + "/" + margin
         + '*100,"0.0")&" points. Per-passenger costs instead reduce the net fare earned on every '
         'seat sold, so they matter most when load factors are high."'),
    ]
    X["insight_first"] = r
    for text in insights:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last_col)
        f = (f'=IFERROR({text[1:]},"Insight not meaningful: Base fare does not exceed the '
             f'variable cost per passenger, or there are no completed flights.")')
        style(ws.cell(row=r, column=1, value=f), font=F_BODY, align=A_WRAP)
        ws.row_dimensions[r].height = 42
        r += 1
    X["insight_last"] = r - 1
    r += 1

    # ---- Validation checks ------------------------------------------------------------
    section_row(ws, r, "6. Sensitivity Validation Checks", last_col)
    r += 1
    on_grid = f"SUMPRODUCT(({lf_axis}={lf})*({fare_axis}={fare}))"
    at_base = f"SUMPRODUCT(({lf_axis}={lf})*({fare_axis}={fare})*{grid})"
    eng_cpf = f"{eng('contrib')}/{eng('completed')}"
    n_cells = SA_LF_ROWS * len(fcols)
    checks = [
        ("Load factor axis within 0%-100%",
         f'IF(AND(MIN({lf_axis})>=0,MAX({lf_axis})<=1),"PASS","FAIL")'),
        ("Matrix rises (or holds) with fare in every row",
         f'IF(SUMPRODUCT(($D${first_r}:$K${last_r}<$C${first_r}:$J${last_r}-0.000001)*1)=0,'
         f'"PASS","FAIL")'),
        ("Matrix rises with load factor wherever fare > variable cost per passenger",
         f'IF(SUMPRODUCT(($C${first_r + 1}:$K${last_r}<$C${first_r}:$K${last_r - 1}-0.000001)'
         f'*({fare_axis}>{vcp}))=0,"PASS","FAIL")'),
        ("Base fare x Base LF cell = Base engine Month 1 contribution / flight",
         f'IF(OR({eng("completed")}=0,{on_grid}<>1),"N/A",'
         f'IF(ABS({at_base}-{eng_cpf})<0.01,"PASS","FAIL"))'),
        ("Flight-level breakeven LF = Base engine Month 1",
         f'IF(OR(NOT(ISNUMBER({b("be_flight")})),{eng("max_margin")}<=0),"N/A",'
         f'IF(ABS({b("be_flight")}-{eng("flight_cost")}/{eng("max_margin")})<0.000001,'
         f'"PASS","FAIL"))'),
        ("Route-level breakeven LF = Base engine Month 1",
         f'IF(OR(NOT(ISNUMBER({b("be_route")})),{eng("max_margin")}<=0),"N/A",'
         f'IF(ABS({b("be_route")}-({eng("flight_cost")}+{eng("fixed")})/{eng("max_margin")})'
         f'<0.000001,"PASS","FAIL"))'),
        ("Contribution per flight at the flight-level breakeven LF = $0",
         f'IF(NOT(ISNUMBER({b("be_flight")})),"N/A",'
         f'IF(ABS({cap}*{b("be_flight")}*({fare}-{vcp})-{vcf})<0.01,"PASS","FAIL"))'),
        ("Route operating profit at the route-level breakeven LF = $0 (Month 1)",
         f'IF(NOT(ISNUMBER({b("be_route")})),"N/A",'
         f'IF(ABS({b("completed")}*({cap}*{b("be_route")}*({fare}-{vcp})-{vcf})-{b("fixed")})'
         f'<0.01,"PASS","FAIL"))'),
        ("Route-level breakeven >= flight-level breakeven",
         f'IF(NOT(ISNUMBER({b("be_gap")})),"N/A",IF({b("be_gap")}>=-0.000001,"PASS","FAIL"))'),
        (f"Legend counts sum to all {n_cells} combinations",
         f'IF(SUM($B${X["cnt_first"]}:$B${X["cnt_last"]})={n_cells},"PASS","FAIL")'),
    ]
    first = r
    for label, f in checks:
        style(ws.cell(row=r, column=1, value=label), font=F_BODY)
        style(ws.cell(row=r, column=2, value=f"={f}"), font=F_BOLD, align=A_CENTER)
        r += 1
    last = r - 1
    rng = f"$B${first}:$B${last}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'],
                                                  fill=FILL_PASS))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'],
                                                  fill=FILL_FAIL))
    style(ws.cell(row=r, column=1, value="SENSITIVITY STATUS"), font=F_BOLD)
    style(ws.cell(row=r, column=2,
                  value=(f'=IF(COUNTIF({rng},"FAIL")=0,"ALL SENSITIVITY CHECKS PASS ("&'
                         f'COUNTIF({rng},"PASS")&" pass, "&COUNTIF({rng},"N/A")&" n/a)",'
                         f'"REVIEW: "&COUNTIF({rng},"FAIL")&" FAILING")')),
          font=F_BOLD, align=A_LEFT)
    ws.conditional_formatting.add(f"B{r}", FormulaRule(formula=[f'LEFT($B${r},3)="ALL"'],
                                                       fill=FILL_PASS))
    ws.conditional_formatting.add(f"B{r}", FormulaRule(formula=[f'LEFT($B${r},3)<>"ALL"'],
                                                       fill=FILL_FAIL))
    X["status"] = r
    X["check_first"], X["check_last"] = first, last

    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 14
    for col in fcols:
        ws.column_dimensions[col].width = 12
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False
    return ws, X


# --------------------------------------------------------------------------
# Executive Dashboard tab (one 16:9 screen; every value links to FM / SC)
# --------------------------------------------------------------------------
DB = "'Executive Dashboard'"
DB_COLS = [get_column_letter(c) for c in range(2, 14)]  # B..M (12 columns)
F_KPI_LABEL = Font(name=FONT_NAME, size=8, bold=True, color=MID_BLUE)
F_KPI_VALUE = Font(name=FONT_NAME, size=20, bold=True, color=NAVY)
F_KPI_SUB = Font(name=FONT_NAME, size=8, color=GREY)
F_KPI_DELTA = Font(name=FONT_NAME, size=8, bold=True, color=NAVY)
FMT_MUSD = '$#,##0.00,,"M";($#,##0.00,,"M");"-"'


def build_dashboard(wb, A, R, S):
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties

    ws = wb.create_sheet("Executive Dashboard", 1)
    D = {}
    fmr = lambda k: f"{FM}!${TOTAL_COL}${R[k]}"  # noqa: E731  FY total in Financial Model
    scr = lambda k, col: f"{SC}!${col}${S[k]}"  # noqa: E731  Scenario Comparison summary
    scenario = f"{AS}!$C${A['scenario']}"
    sc_active = lambda k: f"INDEX({SC}!$C${S[k]}:$E${S[k]},{AS}!$C${A['scenario_idx']})"  # noqa

    ws.column_dimensions["A"].width = 2
    for c in DB_COLS:
        ws.column_dimensions[c].width = 11.5
    ws.column_dimensions["N"].width = 2
    first, last = DB_COLS[0], DB_COLS[-1]

    def band(row, fill, height=None):
        for c in ["A"] + DB_COLS + ["N"]:
            ws[f"{c}{row}"].fill = fill
        if height:
            ws.row_dimensions[row].height = height

    def merged(rng, value, font, fill=None, align=A_LEFT, fmt=None, border=None):
        ws.merge_cells(rng)
        tl = rng.split(":")[0]
        c = ws[tl]
        c.value = value
        style(c, font=font, fill=fill, align=align, fmt=fmt)
        if fill or border:
            for row in ws[rng]:
                for cell in row:
                    if fill:
                        cell.fill = fill
                    if border:
                        cell.border = border
        return c

    # ---- Banner -----------------------------------------------------------
    band(1, FILL_NAVY, 30)
    band(2, FILL_NAVY, 18)
    merged(f"B1:{last}1", "Air Taxi Route Economics  |  Executive Dashboard",
           Font(name=FONT_NAME, size=18, bold=True, color=WHITE), align=A_LEFT)
    merged(f"B2:{last}2",
           (f'="Joby Aviation case study   |   "&{act(A, "route_name")}&"   |   "&'
            f'{act(A, "service_type")}&"   |   12 months "&TEXT({FM}!$C${R["period"]},"mmm yyyy")'
            f'&" - "&TEXT({FM}!${LAST_M}${R["period"]},"mmm yyyy")&"   |   Scenario shown: "'
            f'&{scenario}'),
           Font(name=FONT_NAME, size=10, color="D9E2F3"), align=A_LEFT)
    ws.row_dimensions[3].height = 16
    merged(f"B3:{last}3", "ILLUSTRATIVE DATA ONLY - hypothetical case-study assumptions; not a "
                          "Joby Aviation forecast, guidance or company-reported figure.",
           Font(name=FONT_NAME, size=8, bold=True, italic=True, color="C00000"), align=A_LEFT)
    ws.row_dimensions[4].height = 6

    # ---- KPI cards (rows 5-8) ----------------------------------------------
    is_base = f'{scenario}="Base"'

    def pct_delta(k):
        return (f'=IF({is_base},"Base case shown",IFERROR("vs Base: "&TEXT({fmr(k)}/'
                f'{scr(k, "C")}-1,"+0.0%;-0.0%;0.0%"),"vs Base: n.m."))')

    cards = [
        ("PASSENGER REVENUE", f"={fmr('revenue')}", FMT_MUSD,
         '="USD, 12-month total"', pct_delta("revenue"), "revenue"),
        ("PASSENGERS FLOWN", f"={fmr('pax')}", FMT_NUM,
         '="passengers, 12-month total"', pct_delta("pax"), "pax"),
        ("COMPLETED FLIGHTS", f"={fmr('completed')}", FMT_NUM,
         f'="flights  |  "&TEXT({fmr("comp")},"0.0%")&" completion"',
         pct_delta("completed"), "completed"),
        ("LOAD FACTOR", f"={fmr('lf')}", FMT_PCT,
         '="weighted: passengers / seats"',
         (f'=IF({is_base},"Base case shown",IFERROR("vs Base: "&TEXT(({fmr("lf")}-'
          f'{scr("lf", "C")})*100,"+0.0;-0.0;0.0")&" pts","vs Base: n.m."))'), "lf"),
        ("CONTRIBUTION / FLIGHT", f"={fmr('contrib_flt')}", FMT_USD2,
         '="USD / completed flight, pre-fixed costs"',
         (f'=IF({is_base},"Base case shown",IFERROR("vs Base: "&TEXT({fmr("contrib_flt")}-'
          f'{scr("contrib_flt", "C")},"+$#,##0.00;-$#,##0.00;$0.00"),"vs Base: n.m."))'),
         "contrib_flt"),
        ("ROUTE OPERATING PROFIT", f"={fmr('rop')}", FMT_MUSD,
         f'="after fixed costs  |  "&IFERROR(TEXT({fmr("rom")},"0.0%"),"n.m.")'
         f'&" margin"',
         (f'=IF({is_base},"Base case shown","vs Base: "&TEXT(({fmr("rop")}-{scr("rop", "C")})'
          f'/1000000,"+$0.00;-$0.00;$0.00")&"M")'), "rop"),
    ]
    card_border = Border(left=Side(style="thin", color="BDD7EE"),
                         right=Side(style="thin", color="BDD7EE"),
                         top=Side(style="medium", color=MID_BLUE),
                         bottom=Side(style="thin", color="BDD7EE"))
    for r, h in ((5, 16), (6, 30), (7, 13), (8, 14)):
        ws.row_dimensions[r].height = h
    D["kpi"] = {}
    for i, (label, value, fmt, sub, delta, key) in enumerate(cards):
        c1, c2 = DB_COLS[2 * i], DB_COLS[2 * i + 1]
        for r in range(5, 9):
            for c in (c1, c2):
                ws[f"{c}{r}"].fill = FILL_PALE
        merged(f"{c1}5:{c2}5", label, F_KPI_LABEL, align=Alignment(horizontal="left",
                                                                    vertical="bottom", indent=1))
        merged(f"{c1}6:{c2}6", value, F_KPI_VALUE, fmt=fmt,
               align=Alignment(horizontal="left", vertical="center", indent=1))
        merged(f"{c1}7:{c2}7", sub, F_KPI_SUB, align=Alignment(horizontal="left",
                                                               vertical="center", indent=1))
        merged(f"{c1}8:{c2}8", delta, F_KPI_DELTA, align=Alignment(horizontal="left",
                                                                   vertical="top", indent=1))
        for r in range(5, 9):
            for c in (c1, c2):
                b = ws[f"{c}{r}"].border
                ws[f"{c}{r}"].border = Border(
                    left=card_border.left if c == c1 else b.left,
                    right=card_border.right if c == c2 else b.right,
                    top=card_border.top if r == 5 else b.top,
                    bottom=card_border.bottom if r == 8 else b.bottom)
        D["kpi"][key] = f"{c1}6"
        ws.conditional_formatting.add(f"{c1}6", CellIsRule(operator="lessThan", formula=["0"],
                                                           font=Font(color="C00000")))
    ws.row_dimensions[9].height = 8

    # ---- Section headers (row 10) -----------------------------------------------
    sec_font = Font(name=FONT_NAME, size=10, bold=True, color=NAVY)
    sec_border = Border(bottom=Side(style="medium", color=NAVY))
    ws.row_dimensions[10].height = 18
    merged("B10:H10", "Monthly Passenger Revenue and Route Operating Profit (USD)", sec_font,
           border=sec_border)
    merged(f"I10:{last}10", "Scenario Comparison - 12-Month Totals", sec_font, border=sec_border)

    # ---- Trend chart (links directly to Financial Model rows) --------------------
    fm_ws = wb["Financial Model"]
    cats = Reference(fm_ws, min_col=FIRST_MONTH_COL, max_col=FIRST_MONTH_COL + MONTHS - 1,
                     min_row=R["period"])
    bar = BarChart()
    bar.type = "col"
    bar.add_data(Reference(fm_ws, min_col=2, max_col=FIRST_MONTH_COL + MONTHS - 1,
                           min_row=R["revenue"]), from_rows=True, titles_from_data=False)
    bar.series[0].tx = None
    from openpyxl.chart.series import SeriesLabel
    bar.series[0].tx = SeriesLabel(v="Passenger revenue")
    bar.series[0].graphicalProperties = GraphicalProperties(solidFill="9DC3E6")
    bar.series[0].graphicalProperties.line.noFill = True
    bar.gapWidth = 60
    line = LineChart()
    line.add_data(Reference(fm_ws, min_col=FIRST_MONTH_COL, max_col=FIRST_MONTH_COL + MONTHS - 1,
                            min_row=R["rop"]), from_rows=True, titles_from_data=False)
    line.series[0].tx = SeriesLabel(v="Route operating profit / (loss)")
    line.series[0].graphicalProperties = GraphicalProperties(
        ln=LineProperties(solidFill=NAVY, w=28575))
    line.series[0].marker.symbol = "circle"
    line.series[0].marker.size = 5
    line.series[0].marker.graphicalProperties = GraphicalProperties(solidFill=NAVY)
    line.series[0].smooth = False
    # bar's first data point is column B (unit) - restrict to C..N instead
    bar.series[0].val.numRef.f = (f"{FM}!${FIRST_M}${R['revenue']}:${LAST_M}${R['revenue']}")
    bar.set_categories(cats)
    line.set_categories(cats)
    bar.y_axis.numFmt = '$#,##0,"K"'
    bar.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E7EEF7"))
    bar.x_axis.number_format = "mmm-yy"
    bar.x_axis.delete = False
    bar.x_axis.tickLblPos = "low"  # keep month labels below negative bars/points
    bar.y_axis.delete = False
    bar.y_axis.title = None
    bar.legend.position = "b"
    bar.height = 6.3
    bar.width = 14.6
    bar += line
    ws.add_chart(bar, "B11")
    D["chart_rows"] = (R["revenue"], R["rop"], R["period"])

    # ---- Scenario table (links to Scenario Comparison summary) ------------------
    hdr = 11
    ws.row_dimensions[hdr].height = 16
    merged(f"I{hdr}:J{hdr}", "USD unless noted", Font(name=FONT_NAME, size=8, italic=True,
                                                       color=WHITE), fill=FILL_MID)
    for col, name in zip("KLM", ("Base", "Upside", "Downside")):
        style(ws[f"{col}{hdr}"], font=Font(name=FONT_NAME, size=9, bold=True, color=WHITE),
              fill=FILL_MID, align=A_CENTER)
        ws[f"{col}{hdr}"] = name
    table = [
        ("Passenger revenue", "revenue", FMT_MUSD),
        ("Passengers flown", "pax", FMT_NUM),
        ("Load factor (weighted)", "lf", FMT_PCT),
        ("Average fare ($ / passenger)", "fare", FMT_USD2),
        ("Contribution per flight", "contrib_flt", FMT_USD2),
        ("Contribution profit", "contrib", FMT_MUSD),
        ("Route operating profit", "rop", FMT_MUSD),
        ("Route operating margin", "rom", FMT_PCT),
        ("Route breakeven load factor", "be_route", FMT_PCT),
        ("Profitable months (of 12)", "profit_months", FMT_INT),
    ]
    D["table"] = {}
    t_font = Font(name=FONT_NAME, size=9, color="000000")
    for i, (label, key, fmt) in enumerate(table):
        r = hdr + 1 + i
        ws.row_dimensions[r].height = 16
        fill = FILL_PALE if i % 2 == 0 else PatternFill("solid", fgColor=WHITE)
        bold = key in ("rop", "contrib_flt")
        merged(f"I{r}:J{r}", label, Font(name=FONT_NAME, size=9, bold=bold, color=NAVY), fill=fill)
        for col, sc in zip("KLM", "CDE"):
            style(ws[f"{col}{r}"], font=Font(name=FONT_NAME, size=9, bold=bold, color=NAVY),
                  fmt=fmt, align=A_RIGHT, fill=fill)
            ws[f"{col}{r}"] = f"={scr(key, sc)}"
        D["table"][key] = r
    t_last = hdr + len(table)
    for r in range(hdr + 1, t_last + 1):
        ws[f"{last}{r}"].border = Border(right=Side(style="thin", color="BDD7EE"))
    ws.conditional_formatting.add(f"K{hdr + 1}:M{t_last}", CellIsRule(
        operator="lessThan", formula=["0"], font=Font(color="C00000")))
    # highlight the scenario currently driving the KPI cards
    ws.conditional_formatting.add(f"K{hdr}:M{hdr}", FormulaRule(
        formula=[f"K${hdr}={scenario}"], fill=PatternFill("solid", fgColor=NAVY),
        font=Font(color=WHITE, bold=True, underline="single")))
    r = t_last + 1
    ws.row_dimensions[r].height = 14
    merged(f"I{r}:{last}{r}",
           f'="Navy header = scenario in KPI cards ("&{scenario}&"). Breakeven LF includes '
           f'fixed route costs."', Font(name=FONT_NAME, size=7, italic=True,
                                                     color=GREY))
    # pad rows so the chart area (rows 11-23) has an even height
    for rr in range(r + 1, 24):
        ws.row_dimensions[rr].height = 16
    ws.row_dimensions[24].height = 8

    # ---- Key decisions (rows 25-28) ----------------------------------------------------
    ws.row_dimensions[25].height = 18
    merged(f"B25:{last}25", "Key Decisions for Finance & Operations Leadership "
                            "(scenario shown in KPI cards)", sec_font, border=sec_border)
    fare, lf, be = fmr("fare"), fmr("lf"), fmr("be_route")
    sched, comp, cpf = fmr("sched"), fmr("completed"), fmr("contrib_flt")
    rop_rng = f"{FM}!${FIRST_M}${R['rop']}:${LAST_M}${R['rop']}"
    per_rng = f"{FM}!${FIRST_M}${R['period']}:${LAST_M}${R['period']}"
    decisions = [
        ("PRICING", "Protect the fare; avoid broad discounting until demand is proven.",
         (f'=IFERROR("At a "&TEXT({fare},"$0")&" average fare the route needs a "&TEXT({be},'
          f'"0.0%")&" load factor to cover fixed route costs vs "&TEXT({lf},"0.0%")&" achieved ("'
          f'&TEXT(({lf}-{be})*100,"+0.0;-0.0")&" pts). In the Downside case ("&TEXT('
          f'{scr("fare", "E")},"$0")&" fare, weaker demand) breakeven rises to "&TEXT('
          f'{scr("be_route", "E")},"0.0%")&".","Breakeven load factor is not meaningful in this '
          f'scenario (no passengers flown, or the fare does not exceed variable cost per '
          f'passenger). Revisit fare and demand inputs before setting prices.")')),
        ("CAPACITY", "Prioritise completion rate before adding scheduled frequency.",
         (f'=IFERROR("Completion of "&TEXT({fmr("comp")},"0.0%")&" leaves "&TEXT({sched}-{comp},'
          f'"#,##0")&" of "&TEXT({sched},"#,##0")&" scheduled flights unflown. Each +1 pt of '
          f'completion adds ~"&TEXT({sched}*0.01,"#,##0")&" flights and ~"&TEXT({sched}*0.01*'
          f'{cpf},"$#,##0")&" of annual contribution at "&TEXT({cpf},"$0")&" per flight.",'
          f'"No completed flights in this scenario, so capacity and completion economics are '
          f'not meaningful. Check scheduled flights and completion rate inputs.")')),
        ("ROUTE PROFITABILITY", "Manage the route monthly against its breakeven load factor.",
         (f'=IFERROR("Route operating profit of "&TEXT({fmr("rop")}/1000000,"$0.00;-$0.00")&"M ("'
          f'&TEXT({fmr("rom")},"0.0%")&" margin); "&COUNTIF({rop_rng},">0")&" of 12 months '
          f'profitable"&IFERROR(", first in "&TEXT(INDEX({per_rng},MATCH(1,INDEX(({rop_rng}>0)*1,'
          f'0),0)),"mmm yyyy"),"")&". Fixed route costs need "&TEXT({fmr("fixed")}/{comp},"$0")'
          f'&" of contribution per completed flight.","No completed flights: the route carries "'
          f'&TEXT({fmr("fixed")}/1000000,"$0.00")&"M of fixed route costs with no contribution '
          f'to cover them (route operating profit "&TEXT({fmr("rop")}/1000000,"$0.00;-$0.00")'
          f'&"M).")')),
    ]
    box_cols = [("B", "E"), ("F", "I"), ("J", last)]
    for r, h in ((26, 17), (27, 26), (28, 50)):
        ws.row_dimensions[r].height = h
    D["decisions"] = []
    box_border_l = Side(style="medium", color=MID_BLUE)
    for (c1, c2), (title, decision, evidence) in zip(box_cols, decisions):
        merged(f"{c1}26:{c2}26", title, Font(name=FONT_NAME, size=9, bold=True, color=WHITE),
               fill=FILL_NAVY, align=Alignment(horizontal="left", vertical="center", indent=1))
        merged(f"{c1}27:{c2}27", decision, Font(name=FONT_NAME, size=9, bold=True, color=NAVY),
               fill=FILL_LIGHT, align=Alignment(horizontal="left", vertical="center",
                                                wrap_text=True, indent=1))
        merged(f"{c1}28:{c2}28", evidence, Font(name=FONT_NAME, size=8, color="000000"),
               fill=FILL_PALE, align=Alignment(horizontal="left", vertical="top",
                                               wrap_text=True, indent=1))
        D["decisions"].append(f"{c1}28")
    ws.row_dimensions[29].height = 6

    # ---- Footer -----------------------------------------------------------------
    band(30, FILL_NAVY, 16)
    D["footer"] = f"B30"
    merged(f"B30:{last}30", None, Font(name=FONT_NAME, size=8, color="D9E2F3"))
    D["print_area"] = f"A1:N30"

    # ---- Reconciliation checks (below the screenshot / print area) ---------------
    r = 33
    section_row(ws, r, "Dashboard Reconciliation Checks (outside print area)", 13)
    r += 1
    tol = "0.000001"
    eq = lambda a, b: f"IF(AND(ISNUMBER({a}),ISNUMBER({b})),ABS({a}-{b})<{tol},{a}={b})"  # noqa
    kpi_src = {"revenue": "revenue", "pax": "pax", "completed": "completed", "lf": "lf",
               "contrib_flt": "contrib_flt", "rop": "rop"}
    checks = []
    for key, src in kpi_src.items():
        cell = D["kpi"][key]
        checks.append((f"KPI card {key} = Financial Model FY total", eq(cell, fmr(src))))
    checks.append(("KPI cards = Scenario Comparison active-scenario column (revenue, profit, LF)",
                   "AND(" + ",".join(eq(D["kpi"][k], sc_active(k))
                                     for k in ("revenue", "rop", "lf", "pax")) + ")"))
    checks.append(("Scenario table = Scenario Comparison summary (all cells)",
                   "AND(" + ",".join(eq(f"{col}{D['table'][k]}", scr(k, sc))
                                     for k in D["table"] for col, sc in zip("KLM", "CDE")) + ")"))
    checks.append(("Trend chart revenue months sum to KPI revenue",
                   eq(f"SUM({FM}!${FIRST_M}${R['revenue']}:${LAST_M}${R['revenue']})",
                      D["kpi"]["revenue"])))
    checks.append(("Trend chart profit months sum to KPI route operating profit",
                   eq(f"SUM({rop_rng})", D["kpi"]["rop"])))
    checks.append(("Key decisions text has no errors",
                   "AND(" + ",".join(f"NOT(ISERROR({c}))" for c in D["decisions"]) + ")"))
    first_chk = r
    for label, cond in checks:
        style(ws.cell(row=r, column=2, value=label), font=F_BODY)
        style(ws.cell(row=r, column=10, value=f'=IF({cond},"PASS","FAIL")'), font=F_BOLD,
              align=A_CENTER)
        r += 1
    rng = f"$J${first_chk}:$J${r - 1}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'],
                                                  fill=FILL_PASS))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'],
                                                  fill=FILL_FAIL))
    style(ws.cell(row=r, column=2, value="DASHBOARD STATUS"), font=F_BOLD)
    n = len(checks)
    style(ws.cell(row=r, column=10, value=(f'=IF(COUNTIF({rng},"PASS")={n},"ALL {n} DASHBOARD '
                                           f'CHECKS PASS","REVIEW: "&({n}-COUNTIF({rng},"PASS"))'
                                           f'&" FAILING")')), font=F_BOLD)
    D["status"] = r

    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ws.print_area = D["print_area"]
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.3
    return ws, D


# --------------------------------------------------------------------------
# Checks (bottom of the Financial Model tab)
# --------------------------------------------------------------------------
def build_checks(wb, fm, R, row, last_col, A, S, X, D):
    row += 1
    section_row(fm, row, "E. Model Integrity Checks", last_col)
    row += 1
    for c, v in ((1, "Check"), (3, "Result"), (4, "What it tests")):
        style(fm.cell(row=row, column=c, value=v), font=F_HEADER, fill=FILL_MID)
    for c in range(1, last_col + 1):
        fm.cell(row=row, column=c).fill = FILL_MID
    row += 1

    def mrng(k):
        return f"${FIRST_M}${R[k]}:${LAST_M}${R[k]}"

    def t(k):
        return f"${TOTAL_COL}${R[k]}"

    def arng(k, cols="CE"):
        return f"{AS}!${cols[0]}${A[k]}:${cols[1]}${A[k]}"

    tol = "0.01"
    sc_active = f"INDEX({SC}!$C${{r}}:$E${{r}},{AS}!$C${A['scenario_idx']})"
    checks = [
        ("Scenario selector is valid", f"ISNUMBER({AS}!$C${A['scenario_idx']})",
         "Active scenario name matches Base, Upside or Downside."),
        ("Input load factors within 0%-100% (all scenarios)",
         f"AND(MIN({arng('load_factor')})>=0,MAX({arng('load_factor')})<=1)",
         "Base/Upside/Downside Month 1 load factor inputs."),
        ("Input completion rates within 0%-100% (all scenarios)",
         f"AND(MIN({arng('completion')})>=0,MAX({arng('completion')})<=1)",
         "Base/Upside/Downside completion rate inputs."),
        ("Modelled monthly load factor within 0%-100%",
         f"AND(MIN({mrng('lf')})>=0,MAX({mrng('lf')})<=1)",
         "After demand growth is applied (row is capped at 100%)."),
        ("Modelled monthly completion rate within 0%-100%",
         f"AND(MIN({mrng('comp')})>=0,MAX({mrng('comp')})<=1)", "Active case, every month."),
        ("Capacity, flights per day and operating days are non-negative",
         (f"AND({AS}!$C${A['capacity']}>=0,{AS}!$C${A['flights_per_day']}>=0,"
          f"{AS}!$C${A['op_days']}>=0)"), "Route setup inputs."),
        ("Scheduled and completed flights non-negative; completed <= scheduled",
         (f"AND(MIN({mrng('sched')})>=0,MIN({mrng('completed')})>=0,"
          f"SUMPRODUCT(({mrng('completed')}>{mrng('sched')})*1)=0)"), "Every month."),
        ("Passengers flown <= available seats",
         f"SUMPRODUCT(({mrng('pax')}>{mrng('seats')}+0.000001)*1)=0", "Every month."),
        ("Cost inputs non-negative (all scenarios)",
         f"MIN({AS}!$C${A['var_cost_flight']}:$E${A['fixed_monthly']})>=0",
         "Variable cost per flight, per passenger and fixed costs."),
        ("Modelled costs non-negative",
         f"AND(MIN({mrng('flight_cost')})>=0,MIN({mrng('pax_cost')})>=0,MIN({mrng('fixed')})>=0)",
         "Every month, after cost growth."),
        ("Annual totals equal the sum of the months",
         (f"AND(ABS({t('revenue')}-SUM({mrng('revenue')}))<{tol},"
          f"ABS({t('tvc')}-SUM({mrng('tvc')}))<{tol},ABS({t('rop')}-SUM({mrng('rop')}))<{tol})"),
         "Revenue, total variable costs, route operating profit."),
        ("Revenue reconciles to passengers x average fare, month by month",
         f"ABS({t('revenue')}-SUMPRODUCT({mrng('pax')},{mrng('fare')}))<{tol}",
         "Annual revenue = sum of (monthly passengers x monthly fare)."),
        ("Variable costs reconcile (flight + passenger = total)",
         f"ABS({t('flight_cost')}+{t('pax_cost')}-{t('tvc')})<{tol}",
         "No double counting of variable costs."),
        ("Route P&L reconciles: revenue - variable - fixed = route operating profit",
         f"ABS({t('revenue')}-{t('flight_cost')}-{t('pax_cost')}-{t('fixed')}-{t('rop')})<{tol}",
         "Each cost counted exactly once."),
        ("Contribution per flight x completed flights = contribution profit",
         (f"IF(ISNUMBER({t('contrib_flt')}),ABS({t('contrib_flt')}*{t('completed')}-{t('contrib')})"
          f"<{tol},ABS({t('contrib')})<{tol})"), "Annual (zero flights must mean zero contribution)."),
        ("Cumulative route operating profit = annual total",
         f"ABS({t('cum_rop')}-{t('rop')})<{tol}", "Month 12 cumulative vs. FY total."),
        ("Scenario Comparison (active case) matches Financial Model - revenue",
         f"ABS({sc_active.format(r=S['revenue'])}-{t('revenue')})<{tol}",
         "Independent scenario engine reconciles to the main model."),
        ("Scenario Comparison (active case) matches Financial Model - route operating profit",
         f"ABS({sc_active.format(r=S['rop'])}-{t('rop')})<{tol}",
         "Independent scenario engine reconciles to the main model."),
        ("Scenario Comparison (active case) matches Financial Model - passengers",
         f"ABS({sc_active.format(r=S['pax'])}-{t('pax')})<{tol}",
         "Independent scenario engine reconciles to the main model."),
        ("Sensitivity Analysis checks pass",
         f'LEFT({SA}!$B${X["status"]},3)="ALL"',
         "Status of the validation block on the Sensitivity Analysis tab."),
        ("Executive Dashboard reconciles to source tabs",
         f'LEFT({DB}!$J${D["status"]},3)="ALL"',
         "Status of the reconciliation block below the Executive Dashboard."),
    ]
    first = row
    for label, cond, desc in checks:
        fm.cell(row=row, column=1, value=label).font = F_BODY
        style(fm.cell(row=row, column=3, value=f'=IF({cond},"PASS","FAIL")'), font=F_BOLD,
              align=A_CENTER)
        style(fm.cell(row=row, column=4, value=desc), font=F_NOTE, align=A_LEFT)
        row += 1
    last = row - 1
    n = len(checks)
    rng = f"$C${first}:$C${last}"
    fm.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'],
                                                  fill=FILL_PASS))
    fm.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'],
                                                  fill=FILL_FAIL))
    R["check_first"], R["check_last"] = first, last
    fm.cell(row=row, column=1, value="OVERALL MODEL STATUS").font = F_BOLD
    master = (f'=IF(COUNTIF({rng},"PASS")={n},"ALL {n} CHECKS PASS",'
              f'"REVIEW: "&({n}-COUNTIF({rng},"PASS"))&" CHECK(S) FAILING")')
    fm.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
    style(fm.cell(row=row, column=3, value=master), font=F_BOLD, align=A_CENTER)
    for c in range(1, 7):
        fm.cell(row=row, column=c).border = B_TOTAL
    fm.conditional_formatting.add(f"C{row}", FormulaRule(formula=[f'LEFT($C${row},3)="ALL"'],
                                                         fill=FILL_PASS))
    fm.conditional_formatting.add(f"C{row}", FormulaRule(formula=[f'LEFT($C${row},3)<>"ALL"'],
                                                         fill=FILL_FAIL))
    R["master"] = row

    # Informational flag (not a pass/fail check): demand growth truncated by seat capacity
    row += 2
    fm.cell(row=row, column=1,
            value="Months where demand exceeds seat capacity (load factor capped at 100%)").font = F_BODY
    style(fm.cell(row=row, column=3,
                  value=(f"=SUMPRODUCT(({act(A, 'load_factor')}*{mrng('dem_idx')}>1)*1)")),
          font=F_BOLD, fmt=FMT_INT, align=A_CENTER)
    style(fm.cell(row=row, column=4, value=("Information only. Above zero means the demand growth "
                                            "assumption implies more passengers than seats; the "
                                            "excess is not flown or counted as revenue.")),
          font=F_NOTE, align=A_LEFT)
    fm.conditional_formatting.add(f"C{row}", CellIsRule(operator="greaterThan", formula=["0"],
                                                        fill=PatternFill("solid", fgColor="FFEB9C")))
    R["capped"] = row
    return row, n


# --------------------------------------------------------------------------
# Read Me tab
# --------------------------------------------------------------------------
def build_readme(wb, A, R):
    ws = wb["Sheet"]
    ws.title = "Read Me"
    last_col = 4
    title_block(ws, "Joby Aviation - Air Taxi Route Financial Model (Prototype)",
                "Financial Systems Analyst case study  |  Single illustrative route  |  "
                "12-month monthly model with annual totals", last_col)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110
    ws.column_dimensions["C"].width = 3
    ws.column_dimensions["D"].width = 3

    row = 5
    ws.cell(row=row, column=1, value="Model status").font = F_BOLD
    style(ws.cell(row=row, column=2, value=f"={FM}!$C${R['master']}"), font=Font(
        name=FONT_NAME, size=10, bold=True, color="008000"))
    ws.conditional_formatting.add("B5", FormulaRule(formula=['LEFT($B$5,3)="ALL"'],
                                                    fill=FILL_PASS))
    ws.conditional_formatting.add("B5", FormulaRule(formula=['LEFT($B$5,3)<>"ALL"'],
                                                    fill=FILL_FAIL))
    row += 1
    ws.cell(row=row, column=1, value="Active scenario").font = F_BOLD
    style(ws.cell(row=row, column=2, value=f"={AS}!$C${A['scenario']}"), font=F_LINK)
    row += 1
    ws.cell(row=row, column=1, value="FY route operating profit").font = F_BOLD
    style(ws.cell(row=row, column=2, value=f"={FM}!${TOTAL_COL}${R['rop']}"), font=F_LINK,
          fmt=FMT_USD, align=A_LEFT)
    row += 2

    def block(title, items):
        nonlocal row
        section_row(ws, row, title, last_col)
        row += 1
        for k, v in items:
            style(ws.cell(row=row, column=1, value=k), font=F_BOLD, align=A_WRAP)
            style(ws.cell(row=row, column=2, value=v), font=F_BODY, align=A_WRAP)
            lines = max(1, len(v) // 105 + 1)
            ws.row_dimensions[row].height = 14 * lines + 2
            row += 1
        row += 1

    block("Purpose", [
        ("Objective", "Show how air taxi operating drivers (flights, completion, seats, load factor, "
                      "fare, unit costs) flow through to route-level financial performance, support "
                      "near-term decisions (pricing, scheduling, cost targets), and provide a "
                      "structure that can scale into Joby's commercial reporting infrastructure."),
        ("Important", "ILLUSTRATIVE PROTOTYPE. Every input is a hypothetical assumption for a case "
                      "study. Nothing in this workbook is a Joby Aviation forecast, guidance or "
                      "company-reported figure."),
    ])
    block("Workbook Map", [
        ("1. Read Me", "Purpose, navigation, colour legend, methodology and extension roadmap."),
        ("3. Assumptions", "All editable inputs with units and source/notes. Scenario selector "
                           "(cell C5). Column F 'Active Case' feeds the model."),
        ("4. Financial Model", "12-month route model for the active scenario: driver build, "
                               "operating metrics, route P&L, breakeven analysis and integrity "
                               "checks. FY totals in column O."),
        ("5. Scenario Comparison", "Base / Upside / Downside computed side by side by an "
                                   "independent engine (rows 50+), with variance vs. Base."),
        ("2. Executive Dashboard", "One-screen leadership view (16:9, print area A1:N30): six KPI "
                                "cards, monthly revenue / route operating profit trend, "
                                "Base / Upside / Downside table and three key decisions. All "
                                "values link to Financial Model and Scenario Comparison; "
                                "reconciliation checks sit below the print area."),
        ("6. Sensitivity Analysis", "Base scenario, Month 1: contribution per completed "
                                    "flight across average fare (columns) x load factor (rows), "
                                    "shaded loss / breakeven / profitable. Flight-level vs. "
                                    "route-level breakeven load factor, three business insights "
                                    "and its own validation checks."),

    ])
    block("Colour Legend", [
        ("Blue text, yellow fill", "Editable input (hypothetical, illustrative)."),
        ("Black text", "Formula - do not overwrite."),
        ("Green text", "Link pulled from another tab."),
        ("Navy / light-blue bands", "Headers and section titles."),
        ("PASS / FAIL", "Integrity check results (green / red) at the bottom of Financial Model."),
        ("Executive Dashboard", "Uses presentation colours (navy values) for leadership "
                                "screenshots; every value is still a formula link to Financial "
                                "Model or Scenario Comparison."),
    ])
    block("How To Use", [
        ("Step 1", "On Assumptions, choose the active scenario in C5 (Base, Upside, Downside)."),
        ("Step 2", "Edit blue input cells only. Data validation blocks out-of-range rates and "
                   "negative costs."),
        ("Step 3", "Review Financial Model (active case) and Scenario Comparison (all cases)."),
        ("Step 4", "Confirm the model status above reads 'ALL ... CHECKS PASS'."),
    ])
    block("Methodology", [
        ("Scheduled flights", "Operating days x scheduled flights per operating day."),
        ("Completed flights", "Scheduled flights x flight completion rate (expected value)."),
        ("Available seats", "Completed flights x aircraft passenger seats per flight."),
        ("Passengers flown", "Available seats x passenger load factor (load factor = Month 1 input "
                             "x demand growth index, capped at 100%)."),
        ("Passenger revenue", "Passengers flown x average fare (Month 1 fare x fare growth index)."),
        ("Direct flight costs", "Completed flights x variable operating cost per completed flight "
                                "(x cost growth index). Cancelled flights incur no variable cost."),
        ("Passenger-variable costs", "Passengers flown x variable cost per passenger (x cost growth "
                                     "index)."),
        ("Contribution profit", "Passenger revenue - direct flight costs - passenger-variable "
                                "costs. Flight-level economics before any fixed cost."),
        ("Route operating profit", "Contribution profit - route-specific fixed costs. This is a "
                                   "route-level management measure, NOT EBITDA: corporate "
                                   "overhead, R&D, certification and aircraft ownership costs are "
                                   "excluded."),
        ("Breakeven load factor", "(Direct flight costs [+ fixed route costs]) / (available seats x "
                                  "(fare - variable cost per passenger)). Flight-level version "
                                  "excludes fixed costs; route version includes them. Annual "
                                  "figure uses annual totals, i.e. a uniform load factor across "
                                  "the year."),
        ("Annual rates", "Rates are weighted by their natural base (e.g. load factor = passengers / "
                         "seats; fare = revenue / passengers), never simple averages of months. "
                         "Any ratio with a zero denominator (e.g. margin with no revenue, cost "
                         "per flight with no flights) shows 'n.m.' (not meaningful) rather than 0."),
        ("Capacity cap", "Load factor is capped at 100%. The Financial Model shows how many "
                         "months demand growth exceeds seat capacity (information only)."),
        ("Recalculation", "Workbook is set to fully recalculate on open. No circular references, "
                          "no hard-coded results, no external links."),
    ])
    block("Scaling Roadmap", [
        ("Multiple routes", "Add a Route ID to the Assumptions table (one column block or one row "
                            "set per route), replicate the Financial Model block per route, and "
                            "consolidate with SUMIFS on Route ID."),
        ("Aircraft types / markets", "Move capacity and unit costs to an aircraft-type lookup and "
                                     "fares / demand to a market lookup; route rows reference them "
                                     "via INDEX/MATCH."),
        ("Systems integration", "Replace hypothetical inputs with feeds from scheduling, booking, "
                                "maintenance and ERP systems; keep this calculation layer and the "
                                "integrity checks as the reconciliation control."),
    ])
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"


# --------------------------------------------------------------------------
def build(out_path):
    wb = Workbook()
    _, A = build_assumptions(wb)
    fm, R, row, last_col = build_financial_model(wb, A)
    _, S, E = build_scenarios(wb, A)
    _, X = build_sensitivity(wb, A, E)
    dash, D = build_dashboard(wb, A, R, S)
    master_row, n_checks = build_checks(wb, fm, R, row, last_col, A, S, X, D)
    dash[D["footer"]] = (f'="Sources: Financial Model and Scenario Comparison tabs (formula-linked)'
                         f'   |   Model integrity: "&{FM}!$C${R["master"]}&"   |   Illustrative '
                         f'prototype - not a Joby Aviation forecast"')
    build_readme(wb, A, R)

    # Model status on the Assumptions tab
    asn = wb[AS]
    style(asn[f"C{A['status']}"], font=Font(name=FONT_NAME, size=10, bold=True, color="008000"))
    asn[f"C{A['status']}"] = f"={FM}!$C${R['master']}"

    # Financial Model sheet cosmetics
    fm.column_dimensions["A"].width = 58
    fm.column_dimensions["B"].width = 15
    for c in MONTH_COLS:
        fm.column_dimensions[c].width = 12
    fm.column_dimensions[TOTAL_COL].width = 14
    fm.column_dimensions[BASIS_COL].width = 40
    fm.freeze_panes = "C9"
    fm.sheet_view.showGridLines = False

    # Defined names for future Sensitivity / Dashboard tabs
    for key in ["completion", "fare", "load_factor", "var_cost_flight", "var_cost_pax",
                "fixed_monthly", "fare_growth", "demand_growth", "cost_growth", "capacity",
                "flights_per_day", "op_days"]:
        add_name(wb, f"inp_{key}", f"{AS}!$F${A[key]}")
    add_name(wb, "inp_scenario", f"{AS}!$C${A['scenario']}")
    for key in ["completed", "pax", "revenue", "tvc", "contrib", "contrib_flt", "fixed", "rop",
                "be_flight", "be_route", "lf", "fare"]:
        add_name(wb, f"out_{key}", f"{FM}!${TOTAL_COL}${R[key]}")
    add_name(wb, "out_model_status", f"{FM}!$C${R['master']}")

    # Print setup
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_properties.tabColor = MID_BLUE if ws.title in ("Read Me", "Executive Dashboard") else NAVY

    db = wb["Executive Dashboard"]
    db.page_setup.fitToHeight = 1  # dashboard prints on a single landscape page

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.calcMode = "auto"
    wb.active = 0
    wb.save(out_path)
    return A, R, S, n_checks


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_name(
        "Joby_Financial_Model.xlsx")
    A, R, S, n = build(out)
    print(f"Wrote {out}  ({n} integrity checks)")
