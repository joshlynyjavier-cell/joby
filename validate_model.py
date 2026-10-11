"""
Recalculate Joby_Financial_Model.xlsx headlessly in LibreOffice and validate it.

What it does:
  1. Recalculates a copy of the workbook in LibreOffice and scans every cell
     for Excel errors (#REF!, #DIV/0!, #NAME?, ...).
  2. Recomputes all three scenarios independently in Python and compares the
     results to the workbook's cached formula values.
  3. Switches the scenario selector to Base / Upside / Downside and confirms
     the Financial Model and all integrity checks follow.
  4. Negative tests: feeds invalid inputs (load factor > 100%, completion < 0%)
     and confirms the integrity checks flag them.
  5. With --write-cached, replaces the workbook with the LibreOffice-recalculated
     Base-case copy (so viewers that do not calculate still see values) and
     re-sets Excel's "full calculation on load" flag.

Usage:
    python validate_model.py [Joby_Financial_Model.xlsx] [--write-cached]
"""

import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""

ERRORS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A")
TOL = 0.01


# --------------------------------------------------------------------------
# LibreOffice recalculation
# --------------------------------------------------------------------------
class LibreOffice:
    def __init__(self):
        if not shutil.which("soffice"):
            sys.exit("soffice (LibreOffice) not found on PATH")
        self.tmp = tempfile.TemporaryDirectory(prefix="joby-lo-")
        self.profile = Path(self.tmp.name) / "profile"
        self.env = dict(os.environ, SAL_USE_VCLPLUGIN="svp")
        subprocess.run(["soffice", "--headless", "--terminate_after_init",
                        f"-env:UserInstallation={self.profile.as_uri()}"],
                       env=self.env, capture_output=True, timeout=120)
        macro_dir = self.profile / "user" / "basic" / "Standard"
        if not macro_dir.exists():
            sys.exit("LibreOffice did not create a profile; cannot recalculate")
        (macro_dir / "Module1.xba").write_text(MACRO)

    def recalc(self, path):
        before = Path(path).stat().st_mtime_ns
        subprocess.run(
            ["soffice", "--headless", "--norestore",
             f"-env:UserInstallation={self.profile.as_uri()}",
             "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
             str(Path(path).absolute())],
            env=self.env, capture_output=True, timeout=180, check=True)
        if Path(path).stat().st_mtime_ns == before:
            sys.exit("LibreOffice did not rewrite the file; nothing was recalculated")


def scan_errors(path):
    wb = load_workbook(path, data_only=True)
    found = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and any(e in c.value for e in ERRORS):
                    found.append(f"{ws.title}!{c.coordinate}={c.value}")
    return found


def count_formulas(path):
    wb = load_workbook(path)
    return sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row
               if isinstance(c.value, str) and c.value.startswith("="))


# --------------------------------------------------------------------------
# Locate rows by label (keeps the validator independent of the builder)
# --------------------------------------------------------------------------
def find_row(ws, label, col=1, start=1):
    for r in range(start, ws.max_row + 1):
        v = ws.cell(row=r, column=col).value
        if isinstance(v, str) and v.strip().startswith(label):
            return r
    raise KeyError(f"{ws.title}: row '{label}' not found")


def read_inputs(wb):
    ws = wb["Assumptions"]
    labels = {
        "capacity": "Aircraft passenger seats per flight",
        "fpd": "Scheduled flights per operating day",
        "days": "Operating days per month",
        "completion": "Flight completion rate",
        "fare": "Average fare per passenger",
        "lf": "Passenger load factor",
        "vcf": "Variable operating cost per completed flight",
        "vcp": "Variable cost per passenger",
        "fixed": "Monthly route-specific fixed costs",
        "fg": "Monthly fare growth rate",
        "dg": "Monthly demand growth rate",
        "cg": "Monthly operating cost growth rate",
    }
    rows = {k: find_row(ws, v) for k, v in labels.items()}
    scen = {}
    for name, col in (("Base", 3), ("Upside", 4), ("Downside", 5)):
        d = {}
        for k, r in rows.items():
            v = ws.cell(row=r, column=col).value
            d[k] = ws.cell(row=r, column=3).value if v in (None, "all cases") else v
        scen[name] = d
    return scen, rows


def python_model(p):
    """Independent re-implementation of the route economics."""
    t = dict(sched=0, completed=0, seats=0, pax=0, revenue=0, flight_cost=0, pax_cost=0,
             fixed=0, max_margin=0)
    for m in range(12):
        fare = p["fare"] * (1 + p["fg"]) ** m
        lf = min(1.0, max(0.0, p["lf"] * (1 + p["dg"]) ** m))
        ci = (1 + p["cg"]) ** m
        sched = p["days"] * p["fpd"]
        comp = sched * p["completion"]
        seats = comp * p["capacity"]
        pax = seats * lf
        t["sched"] += sched
        t["completed"] += comp
        t["seats"] += seats
        t["pax"] += pax
        t["revenue"] += pax * fare
        t["flight_cost"] += comp * p["vcf"] * ci
        t["pax_cost"] += pax * p["vcp"] * ci
        t["fixed"] += p["fixed"] * ci
        t["max_margin"] += seats * (fare - p["vcp"] * ci)
    t["tvc"] = t["flight_cost"] + t["pax_cost"]
    t["contrib"] = t["revenue"] - t["tvc"]
    t["rop"] = t["contrib"] - t["fixed"]
    t["contrib_flt"] = t["contrib"] / t["completed"]
    t["be_flight"] = t["flight_cost"] / t["max_margin"]
    t["be_route"] = (t["flight_cost"] + t["fixed"]) / t["max_margin"]
    return t


SC_LABELS = {
    "revenue": "Annual passenger revenue",
    "contrib": "Contribution profit",
    "rop": "Route operating profit",
    "contrib_flt": "Contribution profit per completed flight",
    "be_flight": "Breakeven load factor - flight level",
    "be_route": "Breakeven load factor - route",
    "pax": "Passengers flown",
    "completed": "Completed flights",
}
FM_LABELS = {
    "revenue": "Passenger revenue",
    "contrib": "Contribution profit (flight-level)",
    "rop": "Route operating profit",
    "contrib_flt": "Contribution profit per completed flight",
    "be_flight": "Breakeven load factor - flight level",
    "be_route": "Breakeven load factor - route",
    "pax": "Passengers flown",
    "completed": "Completed flights",
}


def close(a, b, rel=1e-9):
    return isinstance(a, (int, float)) and abs(a - b) <= max(TOL, abs(b) * rel)


def fm_status(wb):
    fm = wb["Financial Model"]
    r = find_row(fm, "OVERALL MODEL STATUS")
    first = find_row(fm, "Check") + 1
    fails = [fm.cell(row=i, column=1).value for i in range(first, r)
             if fm.cell(row=i, column=3).value != "PASS"]
    return fm.cell(row=r, column=3).value, fails, r - first


# --------------------------------------------------------------------------
def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    write_cached = "--write-cached" in sys.argv
    src = Path(args[0] if args else Path(__file__).with_name("Joby_Financial_Model.xlsx"))
    lo = LibreOffice()
    work = Path(lo.tmp.name)
    problems = []

    def variant(name, edit=None):
        p = work / f"{name}.xlsx"
        shutil.copy(src, p)
        if edit:
            wb = load_workbook(p)
            edit(wb)
            wb.save(p)
        lo.recalc(p)
        return p

    print(f"Workbook: {src}")
    print(f"Formulas: {count_formulas(src)}")

    # 1-2. Base recalc, error scan and independent comparison of all scenarios
    base = variant("base")
    errs = scan_errors(base)
    print(f"Formula errors after LibreOffice recalc: {len(errs)}")
    problems += errs[:20]
    wb = load_workbook(base, data_only=True)
    scen, _ = read_inputs(wb)
    sc = wb["Scenario Comparison"]
    start = find_row(sc, "Annual Results")
    print("\nScenario results (workbook vs independent Python):")
    print(f"  {'Scenario':<9}{'Revenue':>14}{'Contribution':>14}{'Route op. profit':>18}"
          f"{'Contrib/flight':>16}{'BE LF route':>13}")
    for name, col in (("Base", 3), ("Upside", 4), ("Downside", 5)):
        py = python_model(scen[name])
        for k, label in SC_LABELS.items():
            v = sc.cell(row=find_row(sc, label, start=start), column=col).value
            if not close(v, py[k]):
                problems.append(f"Scenario Comparison {name} {k}: workbook {v} vs python {py[k]}")
        print(f"  {name:<9}{py['revenue']:>14,.0f}{py['contrib']:>14,.0f}{py['rop']:>18,.0f}"
              f"{py['contrib_flt']:>16,.2f}{py['be_route']:>13.1%}")

    # 3. Selector flips the Financial Model; all checks pass in every case
    print("\nScenario selector test (Financial Model follows selector, checks pass):")
    for name in ("Base", "Upside", "Downside"):
        p = base if name == "Base" else variant(
            name.lower(), lambda w, n=name: w["Assumptions"].__setitem__("C5", n))
        wbv = load_workbook(p, data_only=True)
        fm = wbv["Financial Model"]
        py = python_model(scen[name])
        ok = True
        for k, label in FM_LABELS.items():
            v = fm.cell(row=find_row(fm, label), column=15).value
            if not close(v, py[k]):
                ok = False
                problems.append(f"Financial Model ({name}) {k}: workbook {v} vs python {py[k]}")
        status, fails, n = fm_status(wbv)
        if fails:
            problems.append(f"{name}: failing checks {fails}")
        print(f"  {name:<9} totals match: {'yes' if ok else 'NO'}   status: {status}")

    # 4. Negative tests: invalid inputs must be caught
    print("\nNegative tests (invalid inputs should be flagged):")
    a = load_workbook(src)["Assumptions"]
    lf_row = find_row(a, "Passenger load factor")
    comp_row = find_row(a, "Flight completion rate")
    cost_row = find_row(a, "Variable cost per passenger")
    tests = [
        ("Base load factor = 120%", f"C{lf_row}", 1.2, "Input load factors"),
        ("Downside completion = -10%", f"E{comp_row}", -0.1, "Input completion rates"),
        ("Upside cost per passenger = -$5", f"D{cost_row}", -5, "Cost inputs non-negative"),
    ]
    for i, (desc, cell, val, expect) in enumerate(tests):
        p = variant(f"neg{i}", lambda w, c=cell, v=val: w["Assumptions"].__setitem__(c, v))
        status, fails, _ = fm_status(load_workbook(p, data_only=True))
        caught = any(f.startswith(expect) for f in fails)
        if not caught:
            problems.append(f"Negative test not caught: {desc}")
        print(f"  {desc:<34} -> {status}  ({'caught' if caught else 'NOT CAUGHT'})")

    if write_cached and not problems:
        tmp = work / "final.xlsx"
        with zipfile.ZipFile(base) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "xl/workbook.xml":
                    xml = data.decode("utf-8")
                    xml = re.sub(r"<calcPr\b", '<calcPr fullCalcOnLoad="1"', xml, count=1)
                    data = xml.encode("utf-8")
                zout.writestr(item, data)
        shutil.copy(tmp, src)
        print(f"\nWrote recalculated workbook with cached values to {src} "
              f"(fullCalcOnLoad=1).")

    print("\nRESULT:", "ALL VALIDATIONS PASSED" if not problems else "PROBLEMS FOUND")
    for p in problems:
        print("  -", p)
    sys.exit(1 if problems else 0)


if __name__ == "__main__":
    main()
